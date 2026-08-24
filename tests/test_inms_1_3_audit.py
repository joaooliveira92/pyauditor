"""Caracterização da aba enriquecida do INMS 1.3
(`pyauditor.excel.inms_1_3.write.write_sheet`), acionada por `sintetico.py`
quando o CSV bruto tem as colunas de detalhe exigidas (as mesmas do INMS
1.1). Espelha `test_inms_1_1_audit.py` — mesmo shape `ratio`/
`count_distinct`, trocando "incidentes" por "projetos" e sem
`penalty.base_points` (config real do INMS 1.3 não declara — ver
`configs/_shared/inms-03.yaml`).
"""

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from pyauditor.config.categorias import load_categorias
from pyauditor.excel.sintetico import write_sintetico_workbook
from pyauditor.periodo import PeriodoAfericao

_INMS_03_CONFIG = """\
indicator:
  id: INMS-03
  contractual_id: "INMS 1.3"
  name: Projetos atendidos dentro do prazo

scope:
  contract: "40/2022 - Ministério da Cultura"
  orgao: MinC

source:
  csv: inms-03.csv
  delimiter: ";"
  encoding: utf-8
  period_column: "DataHoraSolicitacao"

quality_gates:
  checks:
    - type: not_null
      column: DataHoraFim
    - type: in_set
      column: "No prazo"
      values: ["S", "N"]

calculation:
  shape: ratio
  aggregation: count_distinct
  numerator_filter:
    column: "No prazo"
    equals: S

target:
  operator: ">="
  value: 100.0

penalty:
  step_points: 200
  step_size_pct: 1.0
"""

_CATEGORIAS_YAML = """\
categorias:
  OPERACAO_N3:
    label: "Operação de Projetos"
    inms:
      "1.3": {mode: grupo_executor, catch_all_contains: "(CIT)"}
"""

# Todas as colunas exigidas por `inms_1_3.write.has_required_columns` (base
# do INMS 1.1) — 2 projetos dentro do prazo, 1 fora.
_INMS_03_RAW_CSV_ENRIQUECIDO = (
    'Nº Solicitacao;Atividades;DataHoraSolicitacao;DataHoraLimite;'
    'DataHoraFim;No prazo;Grupo_executor;TecnicoExecutor\n'
    '1;Implantação;01/06/2026 08:00;10/06/2026 08:00;09/06/2026 08:00;S;'
    'OPERACAO_N3;Fulano\n'
    '2;Migração;01/06/2026 08:00;10/06/2026 08:00;09/06/2026 08:00;S;'
    'OPERACAO_N3;Fulano\n'
    '3;Homologação;01/06/2026 08:00;10/06/2026 08:00;12/06/2026 08:00;N;'
    'OPERACAO_N3;Beltrano\n'
)

_MINIMAL_CSV_SEM_DETALHE = (
    # Tem as colunas mínimas que `measurement_source` exige (period_column,
    # quality gates) — mas falta o restante das colunas de detalhe
    # (`Atividades`, `Nº Solicitacao`, `DataHoraLimite`, `TecnicoExecutor`),
    # então `has_required_columns` continua `False`.
    'DataHoraSolicitacao;DataHoraFim;No prazo;Grupo_executor\n'
    '01/06/2026 08:00;01/06/2026 09:00;S;OPERACAO_N3\n'
)


def _write_fixture(tmp_path: Path) -> tuple[Path, Path]:
    config_dir = tmp_path / 'configs'
    data_dir = tmp_path / 'input' / '2026' / '06'
    config_dir.mkdir(parents=True)
    data_dir.mkdir(parents=True)

    (config_dir / 'inms-03.yaml').write_text(_INMS_03_CONFIG, encoding='utf-8')
    (config_dir / 'categorias.yaml').write_text(
        _CATEGORIAS_YAML, encoding='utf-8'
    )
    (data_dir / 'inms-03.csv').write_text(
        _INMS_03_RAW_CSV_ENRIQUECIDO, encoding='utf-8'
    )

    return config_dir, data_dir


def test_enriched_sheet_is_used_when_raw_csv_has_detail_columns(
    tmp_path: Path,
) -> None:
    config_dir, data_dir = _write_fixture(tmp_path)
    categorias_file = load_categorias(config_dir / 'categorias.yaml')
    output_path = tmp_path / 'sintetico.xlsx'
    periodo = PeriodoAfericao(date(2026, 6, 1), date(2026, 6, 30))

    warnings = write_sintetico_workbook(
        categorias_file, config_dir, data_dir, output_path, periodo=periodo
    )

    assert warnings == []
    wb = load_workbook(output_path)
    sheet = wb['INMS 1.3']

    assert sheet['A1'].value == 'INMS 1.3 – Projetos atendidos dentro do prazo'

    section_bars = {
        cell.value
        for row in sheet.iter_rows(min_col=1, max_col=1)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith('SEÇÃO')
    }
    assert section_bars == {
        'SEÇÃO 1 · IDENTIFICAÇÃO',
        'SEÇÃO 2 · RESUMO EXECUTIVO',
        'SEÇÃO 3 · MEMÓRIA DO CÁLCULO CONSOLIDADO',
        'SEÇÃO 4 · DETALHAMENTO POR GRUPO EXECUTOR',
        'SEÇÃO 5 · SUBTOTAIS POR NÍVEL (informação gerencial)',
        'SEÇÃO 6 · PROJETOS FORA DO PRAZO',
        'SEÇÃO 7 · AUDITORIA DO PRAZO CONTRATUAL',
        'SEÇÃO 8 · TEMPO CORRIDO MÉDIO ATÉ A RESOLUÇÃO',
        'SEÇÃO 9 · PENALIDADE (CÁLCULO PRELIMINAR)',
    }


def test_section_2_kpi_row_counts_two_of_three_within_deadline(
    tmp_path: Path,
) -> None:
    config_dir, data_dir = _write_fixture(tmp_path)
    categorias_file = load_categorias(config_dir / 'categorias.yaml')
    output_path = tmp_path / 'sintetico.xlsx'
    periodo = PeriodoAfericao(date(2026, 6, 1), date(2026, 6, 30))

    write_sintetico_workbook(
        categorias_file, config_dir, data_dir, output_path, periodo=periodo
    )

    wb = load_workbook(output_path)
    sheet = wb['INMS 1.3']

    assert sheet['B13'].value == '=COUNTIF($AP$2:$AP$4,"Sim")'
    assert sheet['C13'].value == ('=COUNTIFS($AP$2:$AP$4,"Sim",$X$2:$X$4,"S")')
    assert sheet['D13'].value == ('=COUNTIFS($AP$2:$AP$4,"Sim",$X$2:$X$4,"N")')
    assert sheet['E13'].value == '=IF(B13=0,"Sem ocorrências",C13/B13)'


def test_section_9_penalty_has_no_base_points(tmp_path: Path) -> None:
    """`configs/_shared/inms-03.yaml` (e o fixture acima) não declaram
    `penalty.base_points` — `Penalty.base_points` (Pydantic) tem default
    `0.0`, então a fórmula de penalidade-base sempre resolve a 0, nunca a um
    valor fixo positivo como o INMS 1.1 (165)."""
    config_dir, data_dir = _write_fixture(tmp_path)
    categorias_file = load_categorias(config_dir / 'categorias.yaml')
    output_path = tmp_path / 'sintetico.xlsx'
    periodo = PeriodoAfericao(date(2026, 6, 1), date(2026, 6, 30))

    write_sintetico_workbook(
        categorias_file, config_dir, data_dir, output_path, periodo=periodo
    )

    wb = load_workbook(output_path)
    sheet = wb['INMS 1.3']

    base_row = next(
        cell.row
        for row in sheet.iter_rows(min_col=1, max_col=1)
        for cell in row
        if cell.value == 'Penalidade-base:'
    )
    assert isinstance(base_row, int)
    formula = sheet.cell(row=base_row, column=2).value
    assert isinstance(formula, str)
    assert ',0,0))' in formula.replace(' ', '')


def test_degrades_to_generic_renderer_when_detail_columns_missing(
    tmp_path: Path,
) -> None:
    """CSV sem as colunas de detalhe não pode quebrar o workbook — cai no
    renderer genérico, como o INMS 1.1/1.2 fazem para fixtures
    minimalistas."""
    config_dir, data_dir = _write_fixture(tmp_path)
    (data_dir / 'inms-03.csv').write_text(
        _MINIMAL_CSV_SEM_DETALHE, encoding='utf-8'
    )
    categorias_file = load_categorias(config_dir / 'categorias.yaml')
    output_path = tmp_path / 'sintetico.xlsx'
    periodo = PeriodoAfericao(date(2026, 6, 1), date(2026, 6, 30))

    warnings = write_sintetico_workbook(
        categorias_file, config_dir, data_dir, output_path, periodo=periodo
    )

    assert warnings == []
    wb = load_workbook(output_path)
    sheet = wb['INMS 1.3']
    assert sheet['A1'].value == 'Categoria'
    section_bars = [
        cell.value
        for row in sheet.iter_rows(min_col=1, max_col=1)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith('SEÇÃO')
    ]
    assert section_bars == []


def test_force_recalc_flags_are_set(tmp_path: Path) -> None:
    config_dir, data_dir = _write_fixture(tmp_path)
    categorias_file = load_categorias(config_dir / 'categorias.yaml')
    output_path = tmp_path / 'sintetico.xlsx'
    periodo = PeriodoAfericao(date(2026, 6, 1), date(2026, 6, 30))

    write_sintetico_workbook(
        categorias_file, config_dir, data_dir, output_path, periodo=periodo
    )

    wb = load_workbook(output_path)
    assert wb.calculation.fullCalcOnLoad is True
    assert wb.calculation.forceFullCalc is True
