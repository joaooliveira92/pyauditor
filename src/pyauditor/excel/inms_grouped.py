"""Aba `INMS_BASE_AGRUPADO` do `relatorio_<competência>_consolidado.xlsx` —
uma visão do `INMS_BASE` da mesma aba com agrupamento nativo de linhas do
Excel (Dados > Agrupar). Todo Código INMS segue o mesmo formato pai/filho:
um `"Consolidado"` (total geral, nível 0) com `"Consolidado - MinC"`/
`"Consolidado - MTur"` (nível 1) como filhos — nunca um órgão pai do outro.
Substitui o pooling por Nível (N1/N2/N3) por um detalhamento real em dois
grupos de INMS:

- Por Grupo executor, nos INMS com essa granularidade em `categorias.yaml`
  (hoje: 1.1, 1.2, 1.3, 1.7 — descoberto em runtime via `_grupo_detail_by_
  inms`, não hardcoded).
- Por ativo/sistema, nos INMS "por ativo" listados em
  `_PRECOMPUTED_BREAKDOWN_CODES` (hoje: 1.4, 1.5, 1.9, 1.10, 1.13, 1.14 —
  todo `precomputed_table` com `numerator_column`/`denominator_column`/
  `name_column` configurados).

Por que não é fabricação:
- Grupo executor: cada linha de detalhe roda
  `SHAPE_REGISTRY[shape].calculate()` (o mesmo motor de `measure`) sobre o
  subconjunto de linhas aceitas (pós quality-gate) de um único
  Grupo_executor — somar essas linhas por categoria reproduz exatamente o
  numerador/denominador já publicado no `INMS_BASE` oficial (conferido
  manualmente).
- Por ativo: cada linha do dataset bruto já É o resultado computado de um
  ativo (`numerator_column`/`denominator_column` do `calculation.
  precomputed_table`); somar essas duas colunas reproduz a mesma aritmética
  que `PrecomputedTableStrategy.calculate` já faz internamente para o
  `result_pct` agregado por órgão (também conferido manualmente contra o
  ROM publicado) — só que exposta por ativo em vez de só o agregado.
- Verbatim reorganizado (`_restructure_verbatim`, os INMS sem nenhum dos
  dois detalhamentos): reusa a linha `"Consolidado"` já publicada no
  `INMS_BASE` quando existe (`with_orgao_consolidation`); quando não existe
  (ex. denominador 0 nos dois órgãos neste mês, ou shape nunca consolidado
  como `precomputed_table`), soma numerador/denominador dos dois órgãos
  quando ambos existem, ou o Resultado calculado direto quando o indicador
  é ponto/contagem sem numerador/denominador (ex. INMS 1.8) — nunca inventa
  uma razão que o indicador não tem.

"Consolidado - {órgão}" é sempre o subtotal/valor daquele órgão;
"Consolidado" é sempre a soma dos dois — aritmética direta sobre números já
corretos, nunca uma fórmula contratual nova.

Grupos executores fora de `categorias.yaml` ("Grupo sob análise de
responsabilidade") ficam de fora do detalhamento por grupo: `measure`/
`report` nunca os soma na apuração oficial, incluí-los aqui infllaria os
subtotais silenciosamente.

`add_inms_agrupado_sheet` é chamada por `cli/consolidate.py` logo depois de
`build_consolidated_workbook` montar o workbook em memória — lê o
`INMS_BASE` que acabou de ser escrito nele (nunca abre nada do disco) e
grava a aba nova ao lado. Se a recomputação falhar (configs/CSV brutos
ausentes, por exemplo num ambiente sem `input/`), `consolidate` degrada com
um aviso e publica o consolidado sem essa aba — nunca bloqueia o artefato
financeiro principal por causa dela.
"""

from __future__ import annotations

from dataclasses import dataclass
from math import isnan
from pathlib import Path
from typing import Final, cast

import openpyxl
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.categoria_filter import (
    GRUPO_EXECUTOR_COLUMN,
    base_config_stem,
    compute_categoria_values,
)
from pyauditor.cli.measure_inputs import resolve_measure_inputs
from pyauditor.codes import format_inms_code, format_inms_code_numeric
from pyauditor.config.categorias import CategoriasFile, GrupoExecutorMode
from pyauditor.config.models import IndicatorConfig, PrecomputedTableCalculation
from pyauditor.config.niveis import NIVEL_BY_CATEGORIA
from pyauditor.config.resolution import per_orgao_paths
from pyauditor.engine.pipeline import measurement_source
from pyauditor.engine.quality_gates import QualityGateRunner
from pyauditor.engine.strategies import SHAPE_REGISTRY
from pyauditor.engine.strategies._numbers import parse_decimal
from pyauditor.engine.strategies._target import meets_target, safe_pct
from pyauditor.excel._style import (
    BODY_FONT,
    HEADER_FILL,
    HEADER_FONT,
    LABEL_FONT,
    CellValue,
)
from pyauditor.excel.inms_base import compliance_margin
from pyauditor.periodo import month_bounds

__all__: Final[tuple[str, ...]] = (
    'GroupedSheetResult',
    'add_inms_agrupado_sheet',
    'compute_glosa_item_detail',
)

_INMS_BASE_SHEET: Final[str] = 'INMS_BASE'
_INMS_BASE_AGRUPADO_SHEET: Final[str] = 'INMS_BASE_AGRUPADO'
_ORGAOS: Final[tuple[str, ...]] = ('MinC', 'MTur')
_SEM_NIVEL: Final[str] = '—'
_AUDIT_REVIEW_LABEL: Final[str] = 'Grupo sob análise de responsabilidade'
_PERCENT_FMT: Final[str] = '0.00"%";(0.00"%");-'
# Mesma regra de `excel/orgao_consolidation.py` (constantes lá são privadas,
# não importáveis) — só grava "Consolidado" para shapes cujo pooling entre
# órgãos já é validado em outro lugar do pipeline.
_CONSOLIDATABLE_SHAPES: Final[frozenset[str]] = frozenset(
    {'ratio', 'segmented_ratio', 'count_difference'}
)
# INMS "por ativo" (`precomputed_table` com `numerator_column`/
# `denominator_column`/`name_column`) com detalhamento por ativo: 1.4/1.5/
# 1.14 (`PER_ASSET_CONTRACTUAL_IDS` em `excel/orgao_consolidation.py`) e
# 1.9/1.10/1.13, que têm exatamente a mesma forma de config (mesmas 3
# colunas), só não eram indicadores "por ativo" na origem — a mecânica de
# `_ativo_detail_by_inms` não distingue os dois casos.
_PRECOMPUTED_BREAKDOWN_CODES: Final[frozenset[str]] = frozenset(
    {'1.4', '1.5', '1.9', '1.10', '1.13', '1.14'}
)

_COLUMNS: Final[tuple[str, ...]] = (
    'Competência', 'Item contratual', 'Serviço', 'Grupo operacional',
    'Código INMS', 'Descrição', 'Órgão', 'Meta mínima ou máxima',
    'Sentido da meta', 'Numerador', 'Denominador', 'Resultado calculado',
    'Unidade', 'Conformidade', 'Diferença para a meta',
)
_COLUMN_WIDTHS: Final[dict[int, int]] = {
    1: 12, 2: 34, 3: 34, 4: 12, 5: 12, 6: 42, 7: 14, 8: 16, 9: 12,
    10: 12, 11: 12, 12: 16, 13: 10, 14: 14, 15: 16,
}
_TOP_BORDER: Final = Border(top=Side(style='thin', color='1F2937'))

# (categoria, nível, grupo executor, numerador, denominador)
type GrupoRow = tuple[str, str, str, float, float]


@dataclass(frozen=True, slots=True)
class _CodeInfo:
    target_value: float
    target_operator: str
    consolidatable: bool


@dataclass(frozen=True, slots=True)
class GroupedSheetResult:
    """Saída de `add_inms_agrupado_sheet` — as contagens que `consolidate`
    reporta ao usuário (a aba já foi escrita direto no workbook recebido)."""

    code_groups: int
    org_subgroups: int
    total_rows: int
    breakdown_codes: tuple[str, ...]


def _compute_for_grupo(
    config: IndicatorConfig,
    id_column: str,
    rows_for_grupo: list[dict[str, str]],
) -> tuple[float, float]:
    """Roda o quality gate + a estratégia oficial (`SHAPE_REGISTRY`) sobre um
    único Grupo_executor — o mesmo caminho de `measure`, só que escopado."""
    gate_report = QualityGateRunner(
        config.quality_gates.checks, id_column=id_column
    ).run(rows_for_grupo)
    strategy = SHAPE_REGISTRY[config.calculation.shape]
    result = strategy.calculate(config, gate_report.accepted)
    numerator, denominator = strategy.pool_numerator_denominator(
        result.memoria
    )
    return numerator or 0.0, denominator or 0.0


def _is_subtotal_label(value: CellValue, *, prefix: str) -> bool:
    return isinstance(value, str) and value.startswith(prefix)


def _rows_for_grupo(
    rows: list[dict[str, str]], grupo: str
) -> list[dict[str, str]]:
    return [r for r in rows if r[GRUPO_EXECUTOR_COLUMN] == grupo]


def _grupo_detail_by_inms(
    competencia: str,
    config_dir: Path,
    data_dir: Path,
    scratch_dir: Path,
) -> tuple[dict[str, dict[str, list[GrupoRow]]], dict[str, _CodeInfo]]:
    """Para cada INMS com `mode: grupo_executor` em `categorias.yaml` (em
    qualquer um dos órgãos), devolve o detalhamento por grupo executor real,
    recomputado pelo motor oficial. `scratch_dir` só existe porque
    `resolve_measure_inputs` exige um destino gravável (`target_dir`) — este
    módulo nunca escreve ROM ali, é só leitura de configs/CSV.
    """
    rows_by_inms: dict[str, dict[str, list[GrupoRow]]] = {}
    info_by_inms: dict[str, _CodeInfo] = {}
    periodo = month_bounds(competencia)

    for orgao in _ORGAOS:
        paths = per_orgao_paths(
            config_dir=config_dir,
            data_dir=data_dir,
            output_dir=scratch_dir,
            orgao=orgao,
        )
        inputs, error = resolve_measure_inputs(
            competencia,
            paths.config_dir,
            paths.data_dir,
            paths.output_dir,
            expected_orgao=orgao,
            manifest=paths.manifest,
        )
        if error is not None or inputs is None:
            raise ValueError(f'{orgao}: {error}')

        configs_by_stem = {
            path.stem: (path, cfg) for path, _hash, cfg in inputs.configs
        }

        for inms_key, entries in inputs.per_inms.items():
            grupo_executor_entries = [
                (ck, e) for ck, e in entries if isinstance(e, GrupoExecutorMode)
            ]
            if not grupo_executor_entries:
                continue
            # `MeasureInputs.categorias_file` é `object | None` (assinatura
            # genérica) — não-None é garantido aqui porque `per_inms` (de
            # onde `entries` veio) só é populado junto com ele, na mesma
            # `_load_categorias`.
            categorias_file = inputs.categorias_file
            if not isinstance(categorias_file, CategoriasFile):
                raise ValueError(
                    f'{orgao}: categorias.yaml não carregado apesar de '
                    'per_inms não vazio'
                )

            stem = base_config_stem(inms_key)
            if stem not in configs_by_stem:
                continue
            config_path, config = configs_by_stem[stem]

            bundle = measurement_source(
                config,
                inputs.competencia_data_dir,
                paths.manifest,
                config_path=config_path,
                periodo=periodo,
                strict=False,
                emit_period_filter_logs=False,
            )
            if GRUPO_EXECUTOR_COLUMN not in bundle.fieldnames:
                continue

            real_values = {row[GRUPO_EXECUTOR_COLUMN] for row in bundle.rows}
            per_categoria_values, outros_values = compute_categoria_values(
                grupo_executor_entries, real_values
            )

            org_rows: list[GrupoRow] = []
            for categoria_key, effective_values in per_categoria_values.items():
                categoria = categorias_file.categorias[categoria_key]
                nivel = NIVEL_BY_CATEGORIA.get(categoria_key, _SEM_NIVEL)
                for grupo in sorted(effective_values):
                    num, den = _compute_for_grupo(
                        config,
                        config.source.id_column,
                        _rows_for_grupo(bundle.rows, grupo),
                    )
                    org_rows.append((categoria.label, nivel, grupo, num, den))

            for grupo in sorted(outros_values):
                num, den = _compute_for_grupo(
                    config,
                    config.source.id_column,
                    _rows_for_grupo(bundle.rows, grupo),
                )
                org_rows.append(
                    (_AUDIT_REVIEW_LABEL, _SEM_NIVEL, grupo, num, den)
                )

            rows_by_inms.setdefault(inms_key, {})[orgao] = org_rows
            if config.target is not None and inms_key not in info_by_inms:
                info_by_inms[inms_key] = _CodeInfo(
                    target_value=config.target.value,
                    target_operator=config.target.operator,
                    consolidatable=config.calculation.shape
                    in _CONSOLIDATABLE_SHAPES,
                )

    return rows_by_inms, info_by_inms


def _ativo_detail_by_inms(
    competencia: str,
    config_dir: Path,
    data_dir: Path,
    scratch_dir: Path,
) -> tuple[dict[str, dict[str, list[GrupoRow]]], dict[str, _CodeInfo]]:
    """Para os INMS "por ativo" em `_PRECOMPUTED_BREAKDOWN_CODES` (shape
    `precomputed_table` com `numerator_column`/`denominator_column`/
    `name_column` configurados), devolve o detalhamento por ativo/sistema:
    cada linha do dataset bruto já É o resultado computado de um ativo —
    somar `numerator_column`/`denominator_column` reproduz exatamente a
    mesma aritmética que `PrecomputedTableStrategy.calculate` já faz
    internamente para o `result_pct` agregado por órgão (conferido
    manualmente: bate com o ROM já publicado), só que exposta por ativo em
    vez de só o agregado — nada de quality-gate/estratégia para recomputar,
    é leitura + soma direta.

    `consolidatable=True` sempre: por pedido explícito, "Consolidado" (o
    pai comum de "Consolidado - MinC"/"Consolidado - MTur") soma os
    numerador/denominador dos dois órgãos — a mesma soma que já gera cada
    subtotal por órgão, só estendida ao outro nível. Diferente do resto do
    pipeline (`excel/orgao_consolidation.py`'s `PER_ASSET_CONTRACTUAL_IDS`),
    que não expõe esse total entre órgãos em nenhum outro lugar por não ter
    validado a fórmula contra uma fonte primária — aqui é só aritmética
    direta sobre números já corretos, não uma fórmula contratual nova.
    """
    rows_by_inms: dict[str, dict[str, list[GrupoRow]]] = {}
    info_by_inms: dict[str, _CodeInfo] = {}
    periodo = month_bounds(competencia)

    for orgao in _ORGAOS:
        paths = per_orgao_paths(
            config_dir=config_dir,
            data_dir=data_dir,
            output_dir=scratch_dir,
            orgao=orgao,
        )
        inputs, error = resolve_measure_inputs(
            competencia,
            paths.config_dir,
            paths.data_dir,
            paths.output_dir,
            expected_orgao=orgao,
            manifest=paths.manifest,
        )
        if error is not None or inputs is None:
            raise ValueError(f'{orgao}: {error}')

        configs_by_stem = {
            path.stem: (path, cfg) for path, _hash, cfg in inputs.configs
        }

        for inms_key in _PRECOMPUTED_BREAKDOWN_CODES:
            stem = base_config_stem(inms_key)
            if stem not in configs_by_stem:
                continue
            config_path, config = configs_by_stem[stem]
            calculation = config.calculation
            if not isinstance(calculation, PrecomputedTableCalculation):
                continue
            if (
                calculation.name_column is None
                or calculation.numerator_column is None
                or calculation.denominator_column is None
            ):
                continue

            bundle = measurement_source(
                config,
                inputs.competencia_data_dir,
                paths.manifest,
                config_path=config_path,
                periodo=periodo,
                strict=False,
                emit_period_filter_logs=False,
            )

            org_rows: list[GrupoRow] = []
            for row in bundle.rows:
                nome = row.get(calculation.name_column, '').strip()
                if not nome:
                    continue
                numerador = parse_decimal(
                    row.get(calculation.numerator_column, '') or ''
                )
                base = parse_decimal(
                    row.get(calculation.denominator_column, '') or ''
                )
                if isnan(numerador) or isnan(base):
                    continue
                org_rows.append(('', _SEM_NIVEL, nome, numerador, base))

            rows_by_inms.setdefault(inms_key, {})[orgao] = org_rows
            if config.target is not None and inms_key not in info_by_inms:
                info_by_inms[inms_key] = _CodeInfo(
                    target_value=config.target.value,
                    target_operator=config.target.operator,
                    consolidatable=True,
                )

    return rows_by_inms, info_by_inms


def _conformidade(
    resultado: float,
    denominador: float,
    target_value: float,
    target_operator: str,
) -> str:
    # Denominador 0 (sem atividade elegível no mês) não é falha de
    # performance — mesma leitura do `RatioStrategy` (0 ocorrências não é
    # 0% de conformidade).
    conforms = denominador == 0 or meets_target(
        resultado, target_operator, target_value
    )
    return 'Conforme' if conforms else 'Não conforme'


def compute_glosa_item_detail(
    competencia: str,
    config_dir: Path,
    data_dir: Path,
    scratch_dir: Path,
) -> dict[tuple[str, str], tuple[str, ...]]:
    """Para cada `(Código INMS, Órgão)` com detalhamento por grupo executor
    ou ativo, devolve os itens que não bateram a meta — a fonte de `Item
    Contratual` da aba GLOSAS (`excel/consolidate/workbook.py::
    build_glosas`). Chaveado pela mesma forma numérica que a GLOSAS já usa
    na coluna `Indicador` (`format_inms_code_numeric`, ex. `"1.02"`).

    Mesma computação de `add_inms_agrupado_sheet`
    (`_grupo_detail_by_inms`/`_ativo_detail_by_inms`), chamada
    separadamente: a GLOSAS é montada por `build_consolidated_workbook`
    antes de `INMS_BASE_AGRUPADO` existir (essa aba só é acrescentada
    depois, em `cli/consolidate.py`), então não há uma aba já escrita para
    ler — o resultado é idêntico ao que aquela aba mostra, só chega mais
    cedo.

    Indicadores sem detalhamento (`whole_indicator` de categoria única)
    simplesmente não aparecem no dict — a GLOSAS deixa `Item Contratual`
    vazio para eles, como já fazia.
    """
    rows_by_inms, info_by_inms = _grupo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    ativo_rows_by_inms, ativo_info_by_inms = _ativo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    rows_by_inms.update(ativo_rows_by_inms)
    info_by_inms.update(ativo_info_by_inms)

    result: dict[tuple[str, str], tuple[str, ...]] = {}
    for inms_key, by_orgao in rows_by_inms.items():
        info = info_by_inms.get(inms_key)
        if info is None:
            continue
        code_key = format_inms_code_numeric(f'INMS {inms_key}')
        for orgao in _ORGAOS:
            org_rows = [
                row for row in by_orgao.get(orgao, [])
                if row[0] != _AUDIT_REVIEW_LABEL
            ]
            failing = tuple(
                grupo
                for _categoria_label, _nivel, grupo, num, den in org_rows
                if _conformidade(
                    round(safe_pct(num, den), 2),
                    den,
                    info.target_value,
                    info.target_operator,
                )
                == 'Não conforme'
            )
            if failing:
                result[code_key, orgao] = failing
    return result


def _consolidado_row(
    competencia: str,
    label: str,
    orgao: str,
    inms_code: str,
    descricao: str | None,
    info: _CodeInfo,
    numerator: float,
    denominator: float,
) -> list[CellValue]:
    resultado = round(safe_pct(numerator, denominator), 2)
    conforme = _conformidade(
        resultado, denominator, info.target_value, info.target_operator
    )
    diff = compliance_margin(resultado, info.target_value, info.target_operator)
    return [
        competencia, label, None, None, inms_code, descricao, orgao,
        info.target_value, info.target_operator, numerator, denominator,
        resultado, '%', conforme, round(diff, 2) if diff is not None else None,
    ]


def _restructure_verbatim(
    competencia: str, inms_code: str, rows: list[list[CellValue]]
) -> list[list[CellValue]] | None:
    """Reorganiza as linhas verbatim de um Código INMS sem detalhamento por
    grupo executor/ativo no mesmo formato pai/filho dos demais: um
    `"Consolidado"` (nível 0) com `"Consolidado - MinC"`/`"Consolidado -
    MTur"` (nível 1) como filhos — só relabela `Item contratual`, nunca
    recalcula Numerador/Denominador/Resultado calculado quando o
    `INMS_BASE` já publica uma linha `"Consolidado"` (reusa verbatim).

    Quando não existe linha `"Consolidado"` publicada (`with_orgao_
    consolidation` não gera uma, seja porque o shape nunca é consolidado —
    `precomputed_table` — seja porque os dois órgãos tiveram denominador 0
    neste mês), soma Numerador/Denominador dos dois órgãos quando ambos
    são numéricos (mesma aritmética do subtotal por órgão, só estendida);
    quando nem isso existe (`precomputed_table` sem numerador/denominador,
    ex. INMS 1.8 — ponto/contagem, não razão), soma o Resultado calculado
    dos dois órgãos diretamente, sem fabricar numerador/denominador.

    Devolve `None` quando as linhas não têm exatamente uma por MinC e uma
    por MTur (formato inesperado) — o chamador mantém o verbatim original
    nesse caso, sem reorganizar."""
    by_orgao: dict[str, list[CellValue]] = {}
    existing_grand: list[CellValue] | None = None
    for row in rows:
        orgao = row[6]
        if orgao == 'Consolidado':
            existing_grand = row
        elif isinstance(orgao, str) and orgao in _ORGAOS:
            if orgao in by_orgao:
                return None
            by_orgao[orgao] = row

    if 'MinC' not in by_orgao or 'MTur' not in by_orgao:
        return None

    def _relabel(row: list[CellValue], label: str) -> list[CellValue]:
        new_row = list(row)
        new_row[1] = label
        return new_row

    children = [
        _relabel(by_orgao[orgao], f'Consolidado - {orgao}') for orgao in _ORGAOS
    ]

    if existing_grand is not None:
        grand = _relabel(list(existing_grand), 'Consolidado')
        return [grand, *children]

    minc_row, mtur_row = by_orgao['MinC'], by_orgao['MTur']
    target_value, target_operator = minc_row[7], minc_row[8]
    if not isinstance(target_value, int | float) or not isinstance(
        target_operator, str
    ):
        return None

    num_minc, den_minc = minc_row[9], minc_row[10]
    num_mtur, den_mtur = mtur_row[9], mtur_row[10]
    if (
        isinstance(num_minc, int | float)
        and isinstance(den_minc, int | float)
        and isinstance(num_mtur, int | float)
        and isinstance(den_mtur, int | float)
    ):
        info = _CodeInfo(target_value, target_operator, consolidatable=True)
        grand = _consolidado_row(
            competencia,
            'Consolidado',
            'Consolidado',
            inms_code,
            minc_row[5] if isinstance(minc_row[5], str) else None,
            info,
            num_minc + num_mtur,
            den_minc + den_mtur,
        )
        return [grand, *children]

    # Sem numerador/denominador (ex. INMS 1.8: ponto/contagem, não razão) —
    # soma o Resultado calculado direto, sem inventar uma razão que o
    # indicador não tem.
    res_minc, res_mtur = minc_row[11], mtur_row[11]
    if not isinstance(res_minc, int | float) or not isinstance(
        res_mtur, int | float
    ):
        return None
    resultado = round(res_minc + res_mtur, 2)
    conforme = _conformidade(resultado, 1.0, target_value, target_operator)
    diff = compliance_margin(resultado, target_value, target_operator)
    grand: list[CellValue] = [
        competencia, 'Consolidado', None, None, inms_code,
        minc_row[5] if isinstance(minc_row[5], str) else None, 'Consolidado',
        target_value, target_operator, None, None, resultado, minc_row[12],
        conforme, round(diff, 2) if diff is not None else None,
    ]
    return [grand, *children]


def _build_breakdown_rows(
    competencia: str,
    rows_by_inms: dict[str, dict[str, list[GrupoRow]]],
    info_by_inms: dict[str, _CodeInfo],
    descricao_by_inms: dict[str, str | None],
) -> dict[str, list[list[CellValue]]]:
    """Monta as linhas finais (detalhe + subtotal por órgão + total geral,
    quando consolidável) de cada INMS com detalhamento por grupo executor."""
    rows_by_code: dict[str, list[list[CellValue]]] = {}

    for inms_key, by_orgao in rows_by_inms.items():
        info = info_by_inms.get(inms_key)
        if info is None:
            continue
        inms_code = format_inms_code(f'INMS {inms_key}')
        descricao = descricao_by_inms.get(inms_key)

        code_rows: list[list[CellValue]] = []
        org_subtotals: list[tuple[str, float, float]] = []

        for orgao in _ORGAOS:
            org_rows = [
                row for row in by_orgao.get(orgao, [])
                if row[0] != _AUDIT_REVIEW_LABEL
            ]
            org_num = sum(n for _, _, _, n, _ in org_rows)
            org_den = sum(d for _, _, _, _, d in org_rows)
            org_subtotals.append((orgao, org_num, org_den))
            if not org_rows:
                continue

            detail_rows: list[list[CellValue]] = []
            for categoria_label, nivel, grupo, num, den in org_rows:
                resultado = round(safe_pct(num, den), 2)
                conforme = _conformidade(
                    resultado, den, info.target_value, info.target_operator
                )
                diff = compliance_margin(
                    resultado, info.target_value, info.target_operator
                )
                detail_rows.append([
                    competencia, grupo, categoria_label, nivel, inms_code,
                    descricao, orgao, info.target_value, info.target_operator,
                    num, den, resultado, '%', conforme,
                    round(diff, 2) if diff is not None else None,
                ])

            code_rows.append(_consolidado_row(
                competencia, f'Consolidado - {orgao}', orgao, inms_code,
                descricao, info, org_num, org_den,
            ))
            code_rows.extend(detail_rows)

        if info.consolidatable:
            grand_num = sum(n for _, n, _ in org_subtotals)
            grand_den = sum(d for _, _, d in org_subtotals)
            code_rows.insert(0, _consolidado_row(
                competencia, 'Consolidado', 'Consolidado', inms_code,
                descricao, info, grand_num, grand_den,
            ))

        if code_rows:
            rows_by_code[inms_key] = code_rows

    return rows_by_code


def _existing_rows_by_code(ws: Worksheet) -> dict[str, list[list[CellValue]]]:
    """Linhas do `INMS_BASE` já publicado, por Código INMS — a fonte
    verbatim para os indicadores sem detalhamento por grupo executor."""
    by_code: dict[str, list[list[CellValue]]] = {}
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 5).value
        if code is None:
            continue
        # `Cell.value` is typed against every openpyxl-representable scalar
        # (Decimal, datetime, ...); `_new_sheet`/`write_row` (`_style.py`)
        # only ever write `CellValue`s, so this narrowing is safe for a
        # workbook this pipeline generated itself.
        row_values = cast(
            'list[CellValue]',
            [ws.cell(r, c).value for c in range(1, 16)],
        )
        by_code.setdefault(str(code), []).append(row_values)
    return by_code


def _write_rows(ws: Worksheet, final_rows: list[list[CellValue]]) -> None:
    for row_idx, values in enumerate(final_rows, start=2):
        label = values[1]
        is_subtotal = _is_subtotal_label(label, prefix='Consolidado')
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row_idx, c, v)
            cell.font = LABEL_FONT if is_subtotal else BODY_FONT
            if c in (8, 12, 15) and isinstance(v, int | float):
                cell.number_format = _PERCENT_FMT
        if is_subtotal:
            for c in range(1, len(_COLUMNS) + 1):
                ws.cell(row_idx, c).border = _TOP_BORDER


def _apply_breakdown_outline(
    ws: Worksheet, r: int, block: list[list[CellValue]]
) -> int:
    """Agrupamento de um bloco com detalhamento.

    Quando existe total geral (rótulo exatamente `"Consolidado"` —
    `_CodeInfo.consolidatable=True`, hoje sempre o caso para os INMS com
    detalhamento), ele é o único nível 0 e cada `"Consolidado - {órgão}"` é
    nível 1, filho dele — um único pai com dois filhos, `Consolidado -
    MinC` e `Consolidado - MTur`.

    Quando não existe (`consolidatable=False` — nenhum código usa isso
    hoje, mas a função continua correta se algum dia usar), os
    `"Consolidado - {órgão}"` viram **irmãos**: todos nível 0, nenhum
    filho do outro — Excel não deixa recolher um irmão de nível 0 sob o
    outro sem um pai comum, e inventar um pai aqui sugeriria uma relação
    entre órgãos que a apuração não valida. Cada um continua com seu
    próprio detalhe recolhido (nível 1 neste caso, em vez de 2).

    Devolve a quantidade de subgrupos `"Consolidado - {órgão}"` criados
    (não conta o total geral, quando existe)."""
    n = len(block)
    has_grand = block[0][1] == 'Consolidado'
    n_subgroups = 0
    offset = 0
    while offset < n:
        header_row = r + offset
        is_grand_row = has_grand and offset == 0
        level = 0 if (is_grand_row or not has_grand) else 1
        ws.row_dimensions[header_row].outlineLevel = level
        ws.row_dimensions[header_row].hidden = has_grand and level != 0
        if not is_grand_row:
            n_subgroups += 1
        offset += 1
        detail_start = offset
        detail_level = level + 1
        while offset < n and not _is_subtotal_label(
            block[offset][1], prefix='Consolidado'
        ):
            ws.row_dimensions[r + offset].outlineLevel = detail_level
            ws.row_dimensions[r + offset].hidden = True
            offset += 1
        if offset > detail_start:
            ws.row_dimensions[header_row].collapsed = True
    if has_grand and n > 1:
        ws.row_dimensions[r].collapsed = True
    return n_subgroups


def _apply_outline(
    ws: Worksheet,
    code_blocks: list[tuple[str, list[list[CellValue]]]],
) -> int:
    """Agrupamento nativo por Código INMS. Todo bloco cuja primeira linha é
    um cabeçalho `"Consolidado"`/`"Consolidado - {órgão}"` (detalhado ou
    verbatim reorganizado por `_restructure_verbatim`) passa por
    `_apply_breakdown_outline` — mesma hierarquia pai/filho em todos os
    casos. Os poucos blocos que `_restructure_verbatim` não conseguiu
    reorganizar (formato inesperado, ex. só uma linha ou nenhuma por
    órgão) caem no fallback verbatim: nível 0 = 1ª linha; nível 1 = demais
    valores distintos de Órgão; nível 2 = repetições. `summaryBelow=False`
    (setado pelo chamador) mantém a linha-resumo acima do seu detalhe.
    Devolve a quantidade de subgrupos "Consolidado - {órgão}" criados."""
    r = 2
    n_org_subgroups = 0
    for _code_full, block in code_blocks:
        n = len(block)
        if _is_subtotal_label(block[0][1], prefix='Consolidado'):
            n_org_subgroups += _apply_breakdown_outline(ws, r, block)
        else:
            orgs_seen: list[CellValue] = []
            for offset in range(n):
                orgao_val = block[offset][6]
                row_i = r + offset
                if orgao_val not in orgs_seen:
                    orgs_seen.append(orgao_val)
                    level = 0 if offset == 0 else 1
                    ws.row_dimensions[row_i].outlineLevel = level
                    ws.row_dimensions[row_i].hidden = level != 0
                    if level == 0:
                        ws.row_dimensions[row_i].collapsed = len(orgs_seen) < n
                    n_org_subgroups += 1
                else:
                    ws.row_dimensions[row_i].outlineLevel = 2
                    ws.row_dimensions[row_i].hidden = True
        r += n
    return n_org_subgroups


def add_inms_agrupado_sheet(
    wb: openpyxl.Workbook,
    competencia: str,
    config_dir: Path,
    data_dir: Path,
    scratch_dir: Path,
) -> GroupedSheetResult:
    """Adiciona a aba `INMS_BASE_AGRUPADO` a `wb` — um workbook consolidado
    já montado em memória por `build_consolidated_workbook`, ainda não
    salvo. Lê o `INMS_BASE` que acabou de ser escrito nele (fonte dos INMS
    sem detalhamento por grupo executor/ativo e da Descrição de cada
    código) e recomputa o detalhamento direto de `config_dir`/`data_dir`.
    Nunca abre nada do disco além disso.

    Raises:
        ValueError: se `INMS_BASE` estiver ausente/vazia em `wb`, ou se
            `resolve_measure_inputs` falhar para um dos órgãos.
    """
    if _INMS_BASE_SHEET not in wb.sheetnames:
        raise ValueError(f'workbook consolidado sem aba {_INMS_BASE_SHEET!r}')
    existing_by_code = _existing_rows_by_code(wb[_INMS_BASE_SHEET])
    if not existing_by_code:
        raise ValueError(f'{_INMS_BASE_SHEET} sem linhas')

    rows_by_inms, info_by_inms = _grupo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    ativo_rows_by_inms, ativo_info_by_inms = _ativo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    # Chaves disjuntas por construção (grupo executor vs. `_PRECOMPUTED_
    # BREAKDOWN_CODES`) — `update` nunca sobrescreve um INMS já presente.
    rows_by_inms.update(ativo_rows_by_inms)
    info_by_inms.update(ativo_info_by_inms)

    descricao_by_inms: dict[str, str | None] = {}
    for inms_key in rows_by_inms:
        code_full = format_inms_code(f'INMS {inms_key}')
        existing_rows = existing_by_code.get(code_full, [])
        descricao = existing_rows[0][5] if existing_rows else None
        descricao_by_inms[inms_key] = (
            descricao if isinstance(descricao, str) else None
        )

    breakdown_rows_by_code = _build_breakdown_rows(
        competencia, rows_by_inms, info_by_inms, descricao_by_inms
    )
    breakdown_by_padded = {
        format_inms_code(f'INMS {k}').replace('INMS ', ''): v
        for k, v in breakdown_rows_by_code.items()
    }

    all_codes = sorted(
        existing_by_code, key=lambda c: float(c.replace('INMS ', ''))
    )
    final_rows: list[list[CellValue]] = []
    code_blocks: list[tuple[str, list[list[CellValue]]]] = []
    for code_full in all_codes:
        inms_key = code_full.replace('INMS ', '')
        if inms_key in breakdown_by_padded:
            block = breakdown_by_padded[inms_key]
        else:
            restructured = _restructure_verbatim(
                competencia, code_full, existing_by_code[code_full]
            )
            block = (
                restructured
                if restructured is not None
                else existing_by_code[code_full]
            )
        code_blocks.append((code_full, block))
        final_rows.extend(block)

    # Logo depois de `INMS_BASE` na ordem das abas — substitui uma execução
    # anterior da mesma competência em vez de duplicar.
    if _INMS_BASE_AGRUPADO_SHEET in wb.sheetnames:
        del wb[_INMS_BASE_AGRUPADO_SHEET]
    insert_at = wb.sheetnames.index(_INMS_BASE_SHEET) + 1
    ws = wb.create_sheet(_INMS_BASE_AGRUPADO_SHEET, index=insert_at)
    ws.sheet_view.showGridLines = False

    for c, name in enumerate(_COLUMNS, start=1):
        cell = ws.cell(1, c, name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for c, width in _COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(c)].width = width

    _write_rows(ws, final_rows)
    n_org_subgroups = _apply_outline(ws, code_blocks)

    ws.row_dimensions[1].outlineLevel = 0
    ws.row_dimensions[1].hidden = False
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_view.showOutlineSymbols = True

    return GroupedSheetResult(
        code_groups=len(code_blocks),
        org_subgroups=n_org_subgroups,
        total_rows=len(final_rows),
        breakdown_codes=tuple(sorted(breakdown_rows_by_code, key=float)),
    )
