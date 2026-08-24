from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from pyauditor.codes import format_inms_code, format_inms_code_numeric
from pyauditor.engine.strategies._target import safe_pct
from pyauditor.excel._style import HEADER_FILL, HEADER_FONT, CellValue

from ._detail import _ativo_detail_by_inms, _grupo_detail_by_inms
from ._rows import (
    _build_breakdown_rows,
    _conformidade,
    _existing_rows_by_code,
    _restructure_verbatim,
)
from ._sheet import _apply_outline, _write_rows
from ._types import (
    _AUDIT_REVIEW_LABEL,
    _COLUMN_WIDTHS,
    _COLUMNS,
    _INMS_BASE_AGRUPADO_SHEET,
    _INMS_BASE_SHEET,
    _ORGAOS,
    GroupedSheetResult,
)


def compute_glosa_item_detail(
    competencia: str,
    config_dir: Path,
    data_dir: Path,
    scratch_dir: Path,
) -> dict[tuple[str, str], tuple[str, ...]]:
    """Para cada `(Código INMS, Órgão)` com detalhamento por grupo executor
    ou ativo, devolve os itens que não bateram a meta — a fonte de `Item
    Contratual` da aba GLOSAS (`excel/consolidate/workbook.py::
    build_glosas`). Chaveado pela mesma forma numérica que a GLOSAS já usa
    na coluna `Indicador` (`format_inms_code_numeric`, ex. `"1.02"`).

    Mesma computação de `add_inms_agrupado_sheet`
    (`_grupo_detail_by_inms`/`_ativo_detail_by_inms`), chamada
    separadamente: a GLOSAS é montada por `build_consolidated_workbook`
    antes de `INMS_BASE_AGRUPADO` existir (essa aba só é acrescentada
    depois, em `cli/consolidate.py`), então não há uma aba já escrita para
    ler — o resultado é idêntico ao que aquela aba mostra, só chega mais
    cedo.

    Indicadores sem detalhamento (`whole_indicator` de categoria única)
    simplesmente não aparecem no dict — a GLOSAS deixa `Item Contratual`
    vazio para eles, como já fazia.
    """
    rows_by_inms, info_by_inms = _grupo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    ativo_rows_by_inms, ativo_info_by_inms = _ativo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    rows_by_inms.update(ativo_rows_by_inms)
    info_by_inms.update(ativo_info_by_inms)

    result: dict[tuple[str, str], tuple[str, ...]] = {}
    for inms_key, by_orgao in rows_by_inms.items():
        info = info_by_inms.get(inms_key)
        if info is None:
            continue
        code_key = format_inms_code_numeric(f'INMS {inms_key}')
        for orgao in _ORGAOS:
            org_rows = [
                row
                for row in by_orgao.get(orgao, [])
                if row[0] != _AUDIT_REVIEW_LABEL
            ]
            failing = tuple(
                grupo
                for _categoria_label, _nivel, grupo, num, den in org_rows
                if _conformidade(
                    round(safe_pct(num, den), 2),
                    den,
                    info.target_value,
                    info.target_operator,
                )
                == 'Não conforme'
            )
            if failing:
                result[code_key, orgao] = failing
    return result


def add_inms_agrupado_sheet(
    wb: openpyxl.Workbook,
    competencia: str,
    config_dir: Path,
    data_dir: Path,
    scratch_dir: Path,
) -> GroupedSheetResult:
    """Adiciona a aba `INMS_BASE_AGRUPADO` a `wb` — um workbook consolidado
    já montado em memória por `build_consolidated_workbook`, ainda não
    salvo. Lê o `INMS_BASE` que acabou de ser escrito nele (fonte dos INMS
    sem detalhamento por grupo executor/ativo e da Descrição de cada
    código) e recomputa o detalhamento direto de `config_dir`/`data_dir`.
    Nunca abre nada do disco além disso.

    Raises:
        ValueError: se `INMS_BASE` estiver ausente/vazia em `wb`, ou se
            `resolve_measure_inputs` falhar para um dos órgãos.
    """
    if _INMS_BASE_SHEET not in wb.sheetnames:
        raise ValueError(f'workbook consolidado sem aba {_INMS_BASE_SHEET!r}')
    existing_by_code = _existing_rows_by_code(wb[_INMS_BASE_SHEET])
    if not existing_by_code:
        raise ValueError(f'{_INMS_BASE_SHEET} sem linhas')

    rows_by_inms, info_by_inms = _grupo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    ativo_rows_by_inms, ativo_info_by_inms = _ativo_detail_by_inms(
        competencia, config_dir, data_dir, scratch_dir
    )
    # Chaves disjuntas por construção (grupo executor vs. `_PRECOMPUTED_
    # BREAKDOWN_CODES`) — `update` nunca sobrescreve um INMS já presente.
    rows_by_inms.update(ativo_rows_by_inms)
    info_by_inms.update(ativo_info_by_inms)

    descricao_by_inms: dict[str, str | None] = {}
    for inms_key in rows_by_inms:
        code_full = format_inms_code(f'INMS {inms_key}')
        existing_rows = existing_by_code.get(code_full, [])
        descricao = existing_rows[0][5] if existing_rows else None
        descricao_by_inms[inms_key] = (
            descricao if isinstance(descricao, str) else None
        )

    breakdown_rows_by_code = _build_breakdown_rows(
        competencia, rows_by_inms, info_by_inms, descricao_by_inms
    )
    breakdown_by_padded = {
        format_inms_code(f'INMS {k}').replace('INMS ', ''): v
        for k, v in breakdown_rows_by_code.items()
    }

    all_codes = sorted(
        existing_by_code, key=lambda c: float(c.replace('INMS ', ''))
    )
    final_rows: list[list[CellValue]] = []
    code_blocks: list[tuple[str, list[list[CellValue]]]] = []
    for code_full in all_codes:
        inms_key = code_full.replace('INMS ', '')
        if inms_key in breakdown_by_padded:
            block = breakdown_by_padded[inms_key]
        else:
            restructured = _restructure_verbatim(
                competencia, code_full, existing_by_code[code_full]
            )
            block = (
                restructured
                if restructured is not None
                else existing_by_code[code_full]
            )
        code_blocks.append((code_full, block))
        final_rows.extend(block)

    # Logo depois de `INMS_BASE` na ordem das abas — substitui uma execução
    # anterior da mesma competência em vez de duplicar.
    if _INMS_BASE_AGRUPADO_SHEET in wb.sheetnames:
        del wb[_INMS_BASE_AGRUPADO_SHEET]
    insert_at = wb.sheetnames.index(_INMS_BASE_SHEET) + 1
    ws = wb.create_sheet(_INMS_BASE_AGRUPADO_SHEET, index=insert_at)
    ws.sheet_view.showGridLines = False

    for c, name in enumerate(_COLUMNS, start=1):
        cell = ws.cell(1, c, name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for c, width in _COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(c)].width = width

    _write_rows(ws, final_rows)
    n_org_subgroups = _apply_outline(ws, code_blocks)

    ws.row_dimensions[1].outlineLevel = 0
    ws.row_dimensions[1].hidden = False
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_view.showOutlineSymbols = True

    return GroupedSheetResult(
        code_groups=len(code_blocks),
        org_subgroups=n_org_subgroups,
        total_rows=len(final_rows),
        breakdown_codes=tuple(sorted(breakdown_rows_by_code, key=float)),
    )
