"""Montagem das linhas finais (`CellValue` por coluna) a partir do
detalhamento recomputado (`_detail.py`) ou do `INMS_BASE` já publicado —
responsabilidade separada de ler configs/CSV (`_detail.py`) e de escrever no
worksheet (`_sheet.py`).
"""

from __future__ import annotations

from typing import cast

from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.codes import format_inms_code
from pyauditor.engine.strategies._target import meets_target, safe_pct
from pyauditor.excel._style import CellValue
from pyauditor.excel.inms_base import compliance_margin

from ._types import _AUDIT_REVIEW_LABEL, _ORGAOS, GrupoRow, _CodeInfo


def _is_subtotal_label(value: CellValue, *, prefix: str) -> bool:
    return isinstance(value, str) and value.startswith(prefix)


def _conformidade(
    resultado: float,
    denominador: float,
    target_value: float,
    target_operator: str,
) -> str:
    # Denominador 0 (sem atividade elegível no mês) não é falha de
    # performance — mesma leitura do `RatioStrategy` (0 ocorrências não é
    # 0% de conformidade).
    conforms = denominador == 0 or meets_target(
        resultado, target_operator, target_value
    )
    return 'Conforme' if conforms else 'Não conforme'


def _consolidado_row(
    competencia: str,
    label: str,
    orgao: str,
    inms_code: str,
    descricao: str | None,
    info: _CodeInfo,
    numerator: float,
    denominator: float,
) -> list[CellValue]:
    resultado = round(safe_pct(numerator, denominator), 2)
    conforme = _conformidade(
        resultado, denominator, info.target_value, info.target_operator
    )
    diff = compliance_margin(resultado, info.target_value, info.target_operator)
    return [
        competencia,
        label,
        None,
        None,
        inms_code,
        descricao,
        orgao,
        info.target_value,
        info.target_operator,
        numerator,
        denominator,
        resultado,
        '%',
        conforme,
        round(diff, 2) if diff is not None else None,
    ]


def _restructure_verbatim(
    competencia: str, inms_code: str, rows: list[list[CellValue]]
) -> list[list[CellValue]] | None:
    """Reorganiza as linhas verbatim de um Código INMS sem detalhamento por
    grupo executor/ativo no mesmo formato pai/filho dos demais: um
    `"Consolidado"` (nível 0) com `"Consolidado - MinC"`/`"Consolidado -
    MTur"` (nível 1) como filhos — só relabela `Item contratual`, nunca
    recalcula Numerador/Denominador/Resultado calculado quando o
    `INMS_BASE` já publica uma linha `"Consolidado"` (reusa verbatim).

    Quando não existe linha `"Consolidado"` publicada (`with_orgao_
    consolidation` não gera uma, seja porque o shape nunca é consolidado —
    `precomputed_table` — seja porque os dois órgãos tiveram denominador 0
    neste mês), soma Numerador/Denominador dos dois órgãos quando ambos
    são numéricos (mesma aritmética do subtotal por órgão, só estendida);
    quando nem isso existe (`precomputed_table` sem numerador/denominador,
    ex. INMS 1.8 — ponto/contagem, não razão), soma o Resultado calculado
    dos dois órgãos diretamente, sem fabricar numerador/denominador.

    Devolve `None` quando as linhas não têm exatamente uma por MinC e uma
    por MTur (formato inesperado) — o chamador mantém o verbatim original
    nesse caso, sem reorganizar."""
    by_orgao: dict[str, list[CellValue]] = {}
    existing_grand: list[CellValue] | None = None
    for row in rows:
        orgao = row[6]
        if orgao == 'Consolidado':
            existing_grand = row
        elif isinstance(orgao, str) and orgao in _ORGAOS:
            if orgao in by_orgao:
                return None
            by_orgao[orgao] = row

    if 'MinC' not in by_orgao or 'MTur' not in by_orgao:
        return None

    def _relabel(row: list[CellValue], label: str) -> list[CellValue]:
        new_row = list(row)
        new_row[1] = label
        return new_row

    children = [
        _relabel(by_orgao[orgao], f'Consolidado - {orgao}') for orgao in _ORGAOS
    ]

    if existing_grand is not None:
        grand = _relabel(list(existing_grand), 'Consolidado')
        return [grand, *children]

    minc_row, mtur_row = by_orgao['MinC'], by_orgao['MTur']
    target_value, target_operator = minc_row[7], minc_row[8]
    if not isinstance(target_value, int | float) or not isinstance(
        target_operator, str
    ):
        return None

    num_minc, den_minc = minc_row[9], minc_row[10]
    num_mtur, den_mtur = mtur_row[9], mtur_row[10]
    if (
        isinstance(num_minc, int | float)
        and isinstance(den_minc, int | float)
        and isinstance(num_mtur, int | float)
        and isinstance(den_mtur, int | float)
    ):
        info = _CodeInfo(target_value, target_operator, consolidatable=True)
        grand = _consolidado_row(
            competencia,
            'Consolidado',
            'Consolidado',
            inms_code,
            minc_row[5] if isinstance(minc_row[5], str) else None,
            info,
            num_minc + num_mtur,
            den_minc + den_mtur,
        )
        return [grand, *children]

    # Sem numerador/denominador (ex. INMS 1.8: ponto/contagem, não razão) —
    # soma o Resultado calculado direto, sem inventar uma razão que o
    # indicador não tem.
    res_minc, res_mtur = minc_row[11], mtur_row[11]
    if not isinstance(res_minc, int | float) or not isinstance(
        res_mtur, int | float
    ):
        return None
    resultado = round(res_minc + res_mtur, 2)
    conforme = _conformidade(resultado, 1.0, target_value, target_operator)
    diff = compliance_margin(resultado, target_value, target_operator)
    grand: list[CellValue] = [
        competencia,
        'Consolidado',
        None,
        None,
        inms_code,
        minc_row[5] if isinstance(minc_row[5], str) else None,
        'Consolidado',
        target_value,
        target_operator,
        None,
        None,
        resultado,
        minc_row[12],
        conforme,
        round(diff, 2) if diff is not None else None,
    ]
    return [grand, *children]


def _build_breakdown_rows(
    competencia: str,
    rows_by_inms: dict[str, dict[str, list[GrupoRow]]],
    info_by_inms: dict[str, _CodeInfo],
    descricao_by_inms: dict[str, str | None],
) -> dict[str, list[list[CellValue]]]:
    """Monta as linhas finais (detalhe + subtotal por órgão + total geral,
    quando consolidável) de cada INMS com detalhamento por grupo executor."""
    rows_by_code: dict[str, list[list[CellValue]]] = {}

    for inms_key, by_orgao in rows_by_inms.items():
        info = info_by_inms.get(inms_key)
        if info is None:
            continue
        inms_code = format_inms_code(f'INMS {inms_key}')
        descricao = descricao_by_inms.get(inms_key)

        code_rows: list[list[CellValue]] = []
        org_subtotals: list[tuple[str, float, float]] = []

        for orgao in _ORGAOS:
            org_rows = [
                row
                for row in by_orgao.get(orgao, [])
                if row[0] != _AUDIT_REVIEW_LABEL
            ]
            org_num = sum(n for _, _, _, n, _ in org_rows)
            org_den = sum(d for _, _, _, _, d in org_rows)
            org_subtotals.append((orgao, org_num, org_den))
            if not org_rows:
                continue

            detail_rows: list[list[CellValue]] = []
            for categoria_label, nivel, grupo, num, den in org_rows:
                resultado = round(safe_pct(num, den), 2)
                conforme = _conformidade(
                    resultado, den, info.target_value, info.target_operator
                )
                diff = compliance_margin(
                    resultado, info.target_value, info.target_operator
                )
                detail_rows.append(
                    [
                        competencia,
                        grupo,
                        categoria_label,
                        nivel,
                        inms_code,
                        descricao,
                        orgao,
                        info.target_value,
                        info.target_operator,
                        num,
                        den,
                        resultado,
                        '%',
                        conforme,
                        round(diff, 2) if diff is not None else None,
                    ]
                )

            code_rows.append(
                _consolidado_row(
                    competencia,
                    f'Consolidado - {orgao}',
                    orgao,
                    inms_code,
                    descricao,
                    info,
                    org_num,
                    org_den,
                )
            )
            code_rows.extend(detail_rows)

        if info.consolidatable:
            grand_num = sum(n for _, n, _ in org_subtotals)
            grand_den = sum(d for _, _, d in org_subtotals)
            code_rows.insert(
                0,
                _consolidado_row(
                    competencia,
                    'Consolidado',
                    'Consolidado',
                    inms_code,
                    descricao,
                    info,
                    grand_num,
                    grand_den,
                ),
            )

        if code_rows:
            rows_by_code[inms_key] = code_rows

    return rows_by_code


def _existing_rows_by_code(ws: Worksheet) -> dict[str, list[list[CellValue]]]:
    """Linhas do `INMS_BASE` já publicado, por Código INMS — a fonte
    verbatim para os indicadores sem detalhamento por grupo executor."""
    by_code: dict[str, list[list[CellValue]]] = {}
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 5).value
        if code is None:
            continue
        # `Cell.value` is typed against every openpyxl-representable scalar
        # (Decimal, datetime, ...); `_new_sheet`/`write_row` (`_style.py`)
        # only ever write `CellValue`s, so this narrowing is safe for a
        # workbook this pipeline generated itself.
        row_values = cast(
            'list[CellValue]',
            [ws.cell(r, c).value for c in range(1, 16)],
        )
        by_code.setdefault(str(code), []).append(row_values)
    return by_code
