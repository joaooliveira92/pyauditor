"""Seções 4–5 (detalhamento por grupo, subtotais) — compartilhadas entre as
abas enriquecidas de INMS; extraídas de `excel/inms_1_1/_sections_4_5.py`
quando o INMS 1.2 ganhou seu próprio renderer enriquecido. Nenhum conteúdo
aqui depende do shape do indicador (ratio vs. segmented_ratio) — a
categoria/grupo executor vem de `categorias.yaml`, não do cálculo do INMS.
"""

from __future__ import annotations

from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common._cells import (
    ColumnRange,
    add_table,
    header_row,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    AG,
    AP,
    AUDIT_REVIEW_LABEL,
    BODY_FONT,
    BORDER,
    DUR,
    GREEN_FILL,
    INCLUIDO_NAO,
    INCLUIDO_SIM,
    LABEL_FONT,
    NIVEL_ORDER_,
    NOTE_FONT,
    ORANGE_FILL,
    PCT2,
    RED_FILL,
    SEM_NIVEL,
    TEAL_FILL,
    UNLOCKED,
    S,
    X,
    Z,
)
from pyauditor.excel._safety import safe_excel_text


def write_section_4_detalhamento(
    sheet: Worksheet,
    *,
    grupo_rows: list[tuple[str, str, str]],
    rng: ColumnRange,
    start_row: int,
    table_name: str,
) -> int:
    """Devolve `next_free_row` — a linha livre após a tabela de grupos,
    usada pela Seção 5 como sua própria linha inicial."""
    section_bar(
        sheet,
        start_row,
        'SEÇÃO 4 · DETALHAMENTO POR GRUPO EXECUTOR',
        last_col=12,
    )
    header_row(
        sheet,
        start_row + 1,
        (
            'Categoria',
            'Nível',
            'Grupo executor',
            'Linhas',
            'Dentro do prazo',
            'Fora do prazo',
            '% (dentro/linhas)',
            'Tempo médio',
            'Incluído no INMS?',
            'Justificativa de exclusão',
            'Documento autorizador',
            'Observação de auditoria',
        ),
        numeric_cols=frozenset({4, 5, 6, 7, 8}),
    )
    first_group_row = start_row + 2
    # Grupos "não previstos" (fora de categorias.yaml) podem ser
    # individualmente desabilitados do cálculo — toggle Sim/Não nesta
    # coluna, lido ao vivo pela coluna de apoio `AP` (Seção 2 em diante).
    # Grupos previstos (categoria conhecida) não são toggleáveis: ficam
    # como texto fixo "Sim", sem validação nem desbloqueio.
    toggle_validation = DataValidation(
        type='list',
        formula1=f'"{INCLUIDO_SIM},{INCLUIDO_NAO}"',
        allow_blank=False,
        errorStyle='stop',
        errorTitle='Valor inválido',
        error=f'Selecione "{INCLUIDO_SIM}" ou "{INCLUIDO_NAO}".',
        showErrorMessage=True,
    )
    # Só registra a validação na aba se algum grupo realmente a usar — um
    # órgão sem nenhum grupo "não previsto" (categoria
    # `AUDIT_REVIEW_LABEL`) nunca chama `.add()`, e registrar mesmo assim
    # grava `<dataValidations count="0" />` no XML: o Excel recusa abrir o
    # arquivo e oferece "reparar", descartando conteúdo.
    has_toggle_target = False
    for offset, (grupo, nivel, categoria) in enumerate(grupo_rows):
        r = first_group_row + offset
        sheet.cell(
            row=r, column=1, value=safe_excel_text(categoria)
        ).font = BODY_FONT
        sheet.cell(
            row=r, column=2, value=safe_excel_text(nivel)
        ).font = BODY_FONT
        sheet.cell(
            row=r, column=3, value=safe_excel_text(grupo)
        ).font = BODY_FONT
        grupo_ref = f'$C${r}'
        linhas_cell = sheet.cell(
            row=r, column=4, value=f'=COUNTIF({rng(S)},{grupo_ref})'
        )
        dentro_cell = sheet.cell(
            row=r,
            column=5,
            value=f'=COUNTIFS({rng(S)},{grupo_ref},{rng(X)},"S")',
        )
        fora_cell = sheet.cell(
            row=r,
            column=6,
            value=f'=COUNTIFS({rng(S)},{grupo_ref},{rng(X)},"N")',
        )
        pct_cell = sheet.cell(
            row=r, column=7, value=f'=IF(D{r}=0,"Sem ocorrências",E{r}/D{r})'
        )
        pct_cell.number_format = PCT2
        tempo_cell = sheet.cell(
            row=r,
            column=8,
            value=f'=IF(D{r}=0,"—",AVERAGEIF({rng(S)},{grupo_ref},{rng(AG)}))',
        )
        tempo_cell.number_format = DUR
        incluido = sheet.cell(row=r, column=9, value=INCLUIDO_SIM)
        is_audit_review = categoria == AUDIT_REVIEW_LABEL
        if is_audit_review:
            just = (
                'Nenhuma exclusão aplicada — ausência de fundamento documental '
                'formal '
                'para exclusão do grupo.'
            )
            obs = (
                'Grupo não previsto em categorias.yaml. Por padrão continua '
                'no consolidado do indicador (denominador = total de '
                'registros abertos no período); para excluí-lo, altere '
                f'"Incluído no INMS?" para "{INCLUIDO_NAO}" e preencha '
                'Justificativa/Documento autorizador ao lado.'
            )
            for cc in (
                linhas_cell,
                dentro_cell,
                fora_cell,
                pct_cell,
                tempo_cell,
                incluido,
            ):
                cc.fill = ORANGE_FILL
            incluido.protection = UNLOCKED
            toggle_validation.add(incluido.coordinate)
            has_toggle_target = True
        else:
            just = 'Não aplicável — nenhuma exclusão aplicada.'
            obs = '—'
        just_cell = sheet.cell(row=r, column=10, value=just)
        doc_cell = sheet.cell(row=r, column=11, value='Não informado')
        obs_cell = sheet.cell(row=r, column=12, value=obs)
        for c in (
            linhas_cell,
            dentro_cell,
            fora_cell,
            pct_cell,
            tempo_cell,
            incluido,
            just_cell,
            doc_cell,
            obs_cell,
        ):
            c.font = BODY_FONT
            c.alignment = Alignment(wrap_text=True, vertical='top')
        for col in range(1, 13):
            sheet.cell(row=r, column=col).border = BORDER
        # Campos de preenchimento manual da auditoria (justificativa/documento
        # autorizador) permanecem editáveis quando a planilha for protegida
        # (ticket 20 / B-03) — os demais campos desta seção são fórmulas.
        just_cell.protection = UNLOCKED
        doc_cell.protection = UNLOCKED
        # Altura maior que o padrão (~15pt) para acomodar o texto quebrado
        # de justificativa/documento/observação (colunas J/K/L) sem cortar
        # visualmente o conteúdo até o usuário redimensionar manualmente.
        sheet.row_dimensions[r].height = 30
    if has_toggle_target:
        sheet.add_data_validation(toggle_validation)
    last_group_row = first_group_row + len(grupo_rows) - 1
    add_table(sheet, table_name, f'A{start_row + 1}:L{last_group_row}')
    sheet.conditional_formatting.add(
        f'I{first_group_row}:I{last_group_row}',
        CellIsRule(
            operator='equal', formula=[f'"{INCLUIDO_NAO}"'], fill=RED_FILL
        ),
    )
    sheet.conditional_formatting.add(
        f'G{first_group_row}:G{last_group_row}',
        ColorScaleRule(
            start_type='min',
            start_color='FECACA',
            mid_type='percentile',
            mid_value=50,
            mid_color='FEF9C3',
            end_type='max',
            end_color='BBF7D0',
        ),
    )
    toggle_note_row = last_group_row + 1
    sheet.merge_cells(f'A{toggle_note_row}:L{toggle_note_row}')
    sheet[f'A{toggle_note_row}'] = (
        f'Apenas grupos não previstos em categorias.yaml (categoria '
        f'"{AUDIT_REVIEW_LABEL}", destacados em laranja) podem ser '
        f'desabilitados — altere "Incluído no INMS?" para '
        f'"{INCLUIDO_NAO}" para excluir o grupo do IAP/IADP consolidado '
        f'(Seção 2) e dos subtotais por nível (Seção 5); registre a '
        f'justificativa e o documento autorizador nas colunas ao lado.'
    )
    sheet[f'A{toggle_note_row}'].font = NOTE_FONT
    sheet[f'A{toggle_note_row}'].alignment = Alignment(
        wrap_text=True, vertical='center'
    )
    return toggle_note_row + 2


def write_section_5_subtotais(
    sheet: Worksheet,
    *,
    rng: ColumnRange,
    start_row: int,
    consolidado_row: int = 13,
) -> int:
    """Devolve `next_free_row` — a linha livre após as verificações
    cruzadas da seção, usada pela Seção 6 como sua própria linha inicial.

    `consolidado_row` — linha da Seção 2 com o IAP/IADP/Fora/Resultado
    consolidado (colunas B/C/D/E) contra o qual a verificação cruzada desta
    seção compara seu próprio total geral; 13 no INMS 1.1 (bloco único de
    KPIs), variável no INMS 1.2 (bloco consolidado após as 3 categorias)."""
    sub_bar_row = start_row
    section_bar(
        sheet,
        sub_bar_row,
        'SEÇÃO 5 · SUBTOTAIS POR NÍVEL (informação gerencial)',
        last_col=6,
    )
    header_row(
        sheet,
        sub_bar_row + 1,
        (
            'Nível / Grupo',
            'Linhas',
            'Dentro do prazo',
            'Fora do prazo',
            '% bruto',
            'Tempo médio',
        ),
        numeric_cols=frozenset({2, 3, 4, 5, 6}),
    )
    nivel_rows = list(
        range(sub_bar_row + 2, sub_bar_row + 2 + len(NIVEL_ORDER_))
    )
    for r, nivel_label in zip(nivel_rows, NIVEL_ORDER_, strict=True):
        sheet.cell(row=r, column=1, value=nivel_label).font = BODY_FONT
        linhas = sheet.cell(
            row=r,
            column=2,
            value=(
                f'=COUNTIFS({rng(Z)},"{nivel_label}",'
                f'{rng(AP)},"{INCLUIDO_SIM}")'
            ),
        )
        dentro = sheet.cell(
            row=r,
            column=3,
            value=(
                f'=COUNTIFS({rng(Z)},"{nivel_label}",{rng(X)},"S",'
                f'{rng(AP)},"{INCLUIDO_SIM}")'
            ),
        )
        fora_c = sheet.cell(
            row=r,
            column=4,
            value=(
                f'=COUNTIFS({rng(Z)},"{nivel_label}",{rng(X)},"N",'
                f'{rng(AP)},"{INCLUIDO_SIM}")'
            ),
        )
        pct = sheet.cell(
            row=r, column=5, value=f'=IF(B{r}=0,"Sem ocorrências",C{r}/B{r})'
        )
        pct.number_format = PCT2
        tempo = sheet.cell(
            row=r,
            column=6,
            value=f'=IF(B{r}=0,"—",AVERAGEIFS({rng(AG)},{rng(Z)},"{nivel_label}"))',
        )
        tempo.number_format = DUR
        for c in (linhas, dentro, fora_c, pct, tempo):
            c.font = BODY_FONT
        for col in range(1, 7):
            sheet.cell(row=r, column=col).border = BORDER

    outr = nivel_rows[-1] + 1
    sheet.cell(
        row=outr, column=1, value=f'Sem nível ({AUDIT_REVIEW_LABEL})'
    ).font = BODY_FONT
    linhas = sheet.cell(
        row=outr,
        column=2,
        value=(f'=COUNTIFS({rng(Z)},"{SEM_NIVEL}",{rng(AP)},"{INCLUIDO_SIM}")'),
    )
    dentro = sheet.cell(
        row=outr,
        column=3,
        value=(
            f'=COUNTIFS({rng(Z)},"{SEM_NIVEL}",{rng(X)},"S",'
            f'{rng(AP)},"{INCLUIDO_SIM}")'
        ),
    )
    fora_c = sheet.cell(
        row=outr,
        column=4,
        value=(
            f'=COUNTIFS({rng(Z)},"{SEM_NIVEL}",{rng(X)},"N",'
            f'{rng(AP)},"{INCLUIDO_SIM}")'
        ),
    )
    pct = sheet.cell(
        row=outr,
        column=5,
        value=f'=IF(B{outr}=0,"Sem ocorrências",C{outr}/B{outr})',
    )
    pct.number_format = PCT2
    tempo = sheet.cell(
        row=outr,
        column=6,
        value=f'=IF(B{outr}=0,"—",AVERAGEIFS({rng(AG)},{rng(Z)},"{SEM_NIVEL}"))',
    )
    tempo.number_format = DUR
    for c in (linhas, dentro, fora_c, pct, tempo):
        c.font = BODY_FONT
        c.fill = ORANGE_FILL
    for col in range(1, 7):
        sheet.cell(row=outr, column=col).border = BORDER

    totr = outr + 1
    sheet.cell(row=totr, column=1, value='Total geral').font = Font(bold=True)
    tl = sheet.cell(row=totr, column=2, value=f'=SUM(B{nivel_rows[0]}:B{outr})')
    td = sheet.cell(row=totr, column=3, value=f'=SUM(C{nivel_rows[0]}:C{outr})')
    tf = sheet.cell(row=totr, column=4, value=f'=SUM(D{nivel_rows[0]}:D{outr})')
    tp = sheet.cell(
        row=totr,
        column=5,
        value=f'=IF(B{totr}=0,"Sem ocorrências",C{totr}/B{totr})',
    )
    tp.number_format = PCT2
    for c in (tl, td, tf, tp):
        c.font = Font(bold=True)
        c.fill = TEAL_FILL
    for col in range(1, 7):
        sheet.cell(row=totr, column=col).border = BORDER

    note_row = totr + 1
    sheet.merge_cells(f'A{note_row}:F{note_row}')
    sheet[f'A{note_row}'] = (
        'O consolidado do indicador é a soma dos numeradores/denominadores '
        'dos níveis '
        '(esta linha), não a média simples dos percentuais de N1/N2/N3.'
    )
    sheet[f'A{note_row}'].font = NOTE_FONT

    check_row = note_row + 1
    sheet[f'A{check_row}'] = 'Verificação cruzada (Seção 5 = Seção 2/3):'
    sheet[f'A{check_row}'].font = LABEL_FONT
    chk = sheet.cell(
        row=check_row,
        column=2,
        value=(
            f'=IF(AND(ISNUMBER(E{totr}),ISNUMBER(E{consolidado_row})),'
            f'IF(ROUND(E{totr},6)=ROUND(E{consolidado_row},6),"OK",'
            f'"DIVERGÊNCIA"),'
            f'IF(E{totr}=E{consolidado_row},"OK","DIVERGÊNCIA"))'
        ),
    )
    chk.font = Font(bold=True)
    sheet.conditional_formatting.add(
        chk.coordinate,
        CellIsRule(operator='equal', formula=['"OK"'], fill=GREEN_FILL),
    )
    sheet.conditional_formatting.add(
        chk.coordinate,
        CellIsRule(operator='equal', formula=['"DIVERGÊNCIA"'], fill=RED_FILL),
    )

    # IAP = dentro do prazo + fora do prazo (ticket 06 / A-04) — se algum
    # valor de "No prazo" escapasse à normalização de `normalize_no_prazo`,
    # essa identidade quebraria silenciosamente (o registro entraria no IAP
    # sem entrar em nenhum dos dois `COUNTIF`); mantida como cinto e
    # suspensório mesmo com a validação na fronteira.
    check2_row = check_row + 1
    sheet[f'A{check2_row}'] = (
        'Verificação (IAP = dentro do prazo + fora do prazo):'
    )
    sheet[f'A{check2_row}'].font = LABEL_FONT
    chk2 = sheet.cell(
        row=check2_row,
        column=2,
        value=(
            f'=IF(B{consolidado_row}=C{consolidado_row}+D{consolidado_row},'
            f'"OK","DIVERGÊNCIA DE CLASSIFICAÇÃO")'
        ),
    )
    chk2.font = Font(bold=True)
    sheet.conditional_formatting.add(
        chk2.coordinate,
        CellIsRule(operator='equal', formula=['"OK"'], fill=GREEN_FILL),
    )
    sheet.conditional_formatting.add(
        chk2.coordinate,
        CellIsRule(
            operator='equal',
            formula=['"DIVERGÊNCIA DE CLASSIFICAÇÃO"'],
            fill=RED_FILL,
        ),
    )
    return check2_row + 2
