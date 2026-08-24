"""Seção 8 (Tempo corrido médio) — compartilhada entre as abas enriquecidas
de INMS; extraída de `excel/inms_1_1/_sections_8_9.py` quando o INMS 1.2
ganhou seu próprio renderer enriquecido. Não depende do shape do cálculo
(ratio vs. segmented_ratio) — é só estatística descritiva sobre a duração
criação→resolução de todas as linhas.
"""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common._cells import (
    ColumnRange,
    label_value,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    AD,
    AF,
    AG,
    BODY_FONT,
    DUR,
    LABEL_FONT,
    NOTE_FONT,
)


def write_section_8_tempo(
    sheet: Worksheet, *, rng: ColumnRange, start_row: int
) -> int:
    """Devolve `next_free_row` — a linha livre após a nota de rodapé da
    seção, usada pela Seção 9 como sua própria linha inicial."""
    s8_bar = start_row
    section_bar(
        sheet,
        s8_bar,
        'SEÇÃO 8 · TEMPO CORRIDO MÉDIO ATÉ A RESOLUÇÃO',
        last_col=6,
    )
    label_value(
        sheet,
        s8_bar + 1,
        'Tempo corrido médio até a resolução (todas as linhas):',
        f'=AVERAGE({rng(AG)})',
        fmt=DUR,
    )
    label_value(
        sheet,
        s8_bar + 2,
        'Mediana do tempo corrido até a resolução:',
        f'=MEDIAN({rng(AG)})',
        fmt=DUR,
    )
    # M-01: `AG` já devolve "" (texto, ignorado por AVERAGE/MEDIAN) para
    # linhas com data ausente/malformada ou encerramento anterior à
    # abertura — mas isso ficava implícito; expõe a contagem de linhas
    # rejeitadas para que a média/mediana acima não pareça cobrir 100% dos
    # registros sem dizer quantos foram excluídos.
    rejeitados_row = s8_bar + 3
    sheet.cell(
        row=rejeitados_row,
        column=1,
        value=(
            'Registros excluídos da média/mediana (data ausente/inválida ou '
            'encerramento antes da abertura):'
        ),
    ).font = LABEL_FONT
    rejeitados_cell = sheet.cell(
        row=rejeitados_row, column=2, value=f'=COUNTIF({rng(AG)},"")'
    )
    rejeitados_cell.font = BODY_FONT
    atraso_row = s8_bar + 4
    sheet.cell(
        row=atraso_row,
        column=1,
        value=(
            'Atraso médio dos registros fora do prazo (vs. limite ITSM, '
            'minutos):'
        ),
    ).font = LABEL_FONT
    # M-02: seleciona pela coluna calculada `AD` ("No prazo (data limite
    # ITSM)"), não pela classificação do fornecedor (`X`) — consistente
    # com o rótulo "vs. limite ITSM"; guarda contra #DIV/0! quando não há
    # nenhum registro fora do prazo por esse critério.
    atraso_cell = sheet.cell(
        row=atraso_row,
        column=2,
        value=(
            f'=IF(COUNTIF({rng(AD)},"N")=0,"Sem '
            f'atrasos",AVERAGEIF({rng(AD)},"N",{rng(AF)}))'
        ),
    )
    atraso_cell.number_format = '0.0'
    atraso_cell.font = BODY_FONT
    note8 = s8_bar + 5
    sheet.merge_cells(f'A{note8}:F{note8}')
    sheet[f'A{note8}'] = (
        'O tempo corrido médio é indicador gerencial complementar — não '
        'substitui a '
        "verificação linha-a-linha do campo 'No prazo' para o cálculo do "
        'indicador.'
    )
    sheet[f'A{note8}'].font = NOTE_FONT
    return note8 + 2
