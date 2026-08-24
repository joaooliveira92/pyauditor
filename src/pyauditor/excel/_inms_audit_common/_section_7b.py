"""Seção 7b (divergência "No prazo" do fornecedor x resultado reproduzido
pela data limite do ITSM) — compartilhada entre as abas enriquecidas de
INMS; extraída de `excel/inms_1_1/_sections_6_7.py` quando o INMS 1.2 ganhou
seu próprio renderer enriquecido. Não depende do prazo contratual bruto (só
de `No prazo` e `DataHoraLimite`/`DataHoraFim`, já reproduzidos na coluna de
apoio `AD` por `write_raw_block_core`), então generaliza para qualquer
indicador que use o mesmo vocabulário S/N.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._datetime import parse_dt
from pyauditor.excel._inms_audit_common._cells import (
    ColumnRange,
    add_table,
    header_row,
)
from pyauditor.excel._inms_audit_common._layout import (
    AD,
    AN,
    AO,
    BODY_FONT,
    BORDER,
    DATETIME_FMT,
    LABEL_FONT,
    NOTE_FONT,
    ORANGE_FILL,
    PCT2,
    R,
    U,
    V,
    W,
    X,
)


def write_section_7b_divergencia_fornecedor(
    sheet: Worksheet,
    *,
    rows: list[dict[str, str]],
    rng: ColumnRange,
    start_row: int,
    table_name: str,
    consolidado_row: int = 13,
) -> int:
    """Sinaliza divergência entre o campo "No prazo" informado pelo
    fornecedor e o resultado reproduzido a partir da data limite registrada
    no próprio ITSM (DataHoraFim ≤ DataHoraLimite) — coluna de apoio AN.
    Diferente de um controle contra o prazo contratual bruto (sensível a
    pausas/suspensão de SLA não modeladas aqui), esta checagem não depende
    de recalcular o prazo do zero: audita a consistência interna entre dois
    campos que o próprio fornecedor já entrega (`No prazo` e
    `DataHoraLimite`), pegando o caso de o campo `No prazo` estar
    desalinhado do limite que o ITSM da contratada já registrou para o
    mesmo chamado."""
    header_row_num = start_row
    sheet.cell(
        row=header_row_num,
        column=1,
        value=(
            'Divergência: "No prazo" informado pelo fornecedor x resultado '
            'reproduzido pela data limite do ITSM (DataHoraFim ≤ '
            'DataHoraLimite)'
        ),
    ).font = LABEL_FONT
    count_row = header_row_num + 1
    sheet.cell(
        row=count_row, column=1, value='Registros divergentes:'
    ).font = BODY_FONT
    div_count = sheet.cell(
        row=count_row, column=2, value=f'=COUNTIF({rng(AN)},"Sim")'
    )
    div_count.font = Font(bold=True)
    sheet.cell(row=count_row, column=3, value='% do total:').font = BODY_FONT
    div_pct = sheet.cell(
        row=count_row,
        column=4,
        value=(
            f'=IF(B{consolidado_row}=0,"Sem ocorrências",'
            f'B{count_row}/B{consolidado_row})'
        ),
    )
    div_pct.number_format = PCT2
    div_pct.font = Font(bold=True)
    div_pct.fill = ORANGE_FILL
    div_count.fill = ORANGE_FILL

    note_row = count_row + 1
    sheet.merge_cells(f'A{note_row}:L{note_row}')
    sheet[f'A{note_row}'] = (
        'Divergência aqui não implica que o fornecedor esteja errado — o '
        'limite ITSM pode ter sido ajustado por pausa/suspensão de SLA não '
        'capturada nesta planilha. Trate como lista de priorização para '
        'checagem manual do histórico do chamado, não como resultado '
        'definitivo.'
    )
    sheet[f'A{note_row}'].font = Font(
        name='Arial', size=10, bold=True, color='9A3412'
    )
    sheet[f'A{note_row}'].fill = ORANGE_FILL
    sheet[f'A{note_row}'].alignment = Alignment(
        wrap_text=True, vertical='center'
    )
    sheet.row_dimensions[note_row].height = 40

    divergentes_count = 0
    for row in rows:
        fim_raw = row.get('DataHoraFim', '')
        limite_raw = row.get('DataHoraLimite', '')
        fim = parse_dt(fim_raw).value
        limite = parse_dt(limite_raw).value
        if fim is None or limite is None:
            continue
        no_prazo_itsm = 'S' if fim <= limite else 'N'
        if row.get('No prazo', '') != no_prazo_itsm:
            divergentes_count += 1
    sample_size = min(divergentes_count, 15)

    sample_header_row = note_row + 2
    if sample_size == 0:
        sheet.merge_cells(f'A{sample_header_row}:F{sample_header_row}')
        sheet[f'A{sample_header_row}'] = (
            'Nenhuma divergência entre "No prazo" do fornecedor e o '
            'resultado reproduzido pela data limite do ITSM encontrada no '
            'período.'
        )
        sheet[f'A{sample_header_row}'].font = NOTE_FONT
        return sample_header_row + 2

    sheet.cell(
        row=sample_header_row,
        column=1,
        value=(
            f'Amostra de registros divergentes ({sample_size} primeiros, '
            f'ordenados por ocorrência)'
        ),
    ).font = LABEL_FONT
    header_row(
        sheet,
        sample_header_row + 1,
        (
            'Nº solicitação',
            'No prazo (fornecedor)',
            'No prazo (data limite ITSM)',
            'Abertura',
            'Limite (ITSM)',
            'Encerramento',
        ),
        numeric_cols=frozenset({4, 5, 6}),
    )
    sample_first = sample_header_row + 2
    ao_range = rng(AO)
    for n in range(1, sample_size + 1):
        r = sample_first + n - 1
        match_expr = f'MATCH({n},{ao_range},0)'
        c1 = sheet.cell(
            row=r, column=1, value=f'=IFERROR(INDEX({rng(R)},{match_expr}),"")'
        )
        c2 = sheet.cell(
            row=r, column=2, value=f'=IFERROR(INDEX({rng(X)},{match_expr}),"")'
        )
        c3 = sheet.cell(
            row=r,
            column=3,
            value=f'=IFERROR(INDEX({rng(AD)},{match_expr}),"")',
        )
        c4 = sheet.cell(
            row=r, column=4, value=f'=IFERROR(INDEX({rng(U)},{match_expr}),"")'
        )
        c4.number_format = DATETIME_FMT
        c5 = sheet.cell(
            row=r, column=5, value=f'=IFERROR(INDEX({rng(V)},{match_expr}),"")'
        )
        c5.number_format = DATETIME_FMT
        c6 = sheet.cell(
            row=r, column=6, value=f'=IFERROR(INDEX({rng(W)},{match_expr}),"")'
        )
        c6.number_format = DATETIME_FMT
        for c in (c1, c2, c3, c4, c5, c6):
            c.font = BODY_FONT
            c.border = BORDER
    sample_last = sample_first + sample_size - 1
    add_table(sheet, table_name, f'A{sample_header_row + 1}:F{sample_last}')
    sample_note_row = sample_last + 1
    sheet.merge_cells(f'A{sample_note_row}:F{sample_note_row}')
    sheet[f'A{sample_note_row}'] = (
        f'=CONCATENATE("Amostra limitada às {sample_size} primeiras '
        f'ocorrências de ",'
        f'B{count_row}," registros divergentes — colunas de apoio desta '
        f'aba (coluna AN) permitem reproduzir a lista completa.")'
    )
    sheet[f'A{sample_note_row}'].font = NOTE_FONT
    return sample_note_row + 2
