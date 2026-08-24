"""Constantes de leiaute compartilhadas entre as abas enriquecidas de INMS
(hoje 1.1 e 1.2) — extraídas de `excel/inms_1_1/_layout.py` quando o INMS 1.2
ganhou seu próprio renderer enriquecido, para as duas abas pararem de
duplicar estilo/colunas de apoio/vocabulário de domínio.
"""

from __future__ import annotations

from typing import Final

from openpyxl.styles import Border, Font, PatternFill, Protection, Side

from pyauditor.config.niveis import NIVEL_BY_CATEGORIA, NIVEL_ORDER

# Colunas do CSV bruto exigidas para o tratamento enriquecido, além das já
# exigidas pelo shape genérico (`No prazo`, `DataHoraSolicitacao`,
# `DataHoraFim`, `Grupo_executor`) — nome idêntico nos CSVs de INMS 1.1 e 1.2.
NUM_SOLICITACAO_COLUMN: Final[str] = 'Nº Solicitacao'
ATIVIDADE_COLUMN: Final[str] = 'Atividades'
DATA_SOLICITACAO_COLUMN: Final[str] = 'DataHoraSolicitacao'
DATA_LIMITE_COLUMN: Final[str] = 'DataHoraLimite'
DATA_FIM_COLUMN: Final[str] = 'DataHoraFim'
NO_PRAZO_COLUMN: Final[str] = 'No prazo'
TECNICO_COLUMN: Final[str] = 'TecnicoExecutor'

NIVEL_BY_CATEGORIA_ = NIVEL_BY_CATEGORIA
NIVEL_ORDER_ = NIVEL_ORDER
AUDIT_REVIEW_LABEL: Final[str] = 'Grupo sob análise de responsabilidade'
# Sentinela para grupos sem Nível (categoria "outros") — não usar "" como
# valor de célula: o Excel/openpyxl trata célula com "" como vazia, e
# VLOOKUP contra uma célula vazia devolve 0 (numérico), não "" (texto),
# quebrando os filtros COUNTIFS por Nível na Seção 5.
SEM_NIVEL: Final[str] = '—'

DATETIME_FMT: Final[str] = 'dd/mm/yyyy hh:mm'
DATE_FMT: Final[str] = 'dd/mm/yyyy'
PCT2: Final[str] = '0.00%'
PCT4: Final[str] = '0.0000%'
DUR: Final[str] = '[h]:mm'

TITLE_FONT: Final = Font(name='Arial', size=14, bold=True, color='1F2937')
SECTION_FONT: Final = Font(name='Arial', size=11, bold=True, color='FFFFFF')
SECTION_FILL: Final = PatternFill('solid', fgColor='1F2937')
HEADER_FONT: Final = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HEADER_FILL: Final = PatternFill('solid', fgColor='374151')
BODY_FONT: Final = Font(name='Arial', size=10)
LABEL_FONT: Final = Font(name='Arial', size=10, bold=True)
NOTE_FONT: Final = Font(name='Arial', size=9, italic=True, color='6B7280')
GRAY_FILL: Final = PatternFill('solid', fgColor='E5E7EB')
TEAL_FILL: Final = PatternFill('solid', fgColor='CCFBF1')
GREEN_FILL: Final = PatternFill('solid', fgColor='BBF7D0')
RED_FILL: Final = PatternFill('solid', fgColor='FECACA')
ORANGE_FILL: Final = PatternFill('solid', fgColor='FED7AA')
BORDER: Final = Border(bottom=Side(style='thin', color='D1D5DB'))
# Campos de preenchimento manual (justificativa/documento/evidência) que
# devem continuar editáveis mesmo com a planilha protegida (ticket 20 / B-03).
UNLOCKED: Final = Protection(locked=False)

# Colunas de apoio (dados brutos) — far à direita do conteúdo visível. Mesmo
# leiaute para as duas abas (1.1 e 1.2): a extensão específica de INMS 1.1
# (controle contratual bruto de N horas corridas, colunas AB/AC/AE/AI) só é
# escrita/lida pela própria aba 1.1 — para o INMS 1.2 essas colunas
# simplesmente não são preenchidas (a Seção 7 daquela aba não as referencia),
# em vez de renumerar o restante do bloco de apoio por conta de um
# subconjunto que não se aplica a todo indicador.
(
    R,
    S,
    T,
    U,
    V,
    W,
    X,
    Y,
    Z,
    AA,
    AB,
    AC,
    AD,
    AE,
    AF,
    AG,
    AH,
    AI,
) = range(18, 36)
AJ = 36
AK, AL, AM = 37, 38, 39
AN, AO = 40, 41
# Flag por linha: grupo do incidente/requisição está habilitado para o
# cálculo do indicador — lookup ao vivo contra a coluna "Incluído no INMS?"
# da Seção 4, então reage ao toggle Sim/Não sem precisar reexecutar o
# pipeline.
AP = 42
# Mapa auxiliar (uma linha por grupo, paralelo a `AK`): referência direta à
# célula "Incluído no INMS?" do grupo na Seção 4 — é o que `AP` consulta via
# INDEX/MATCH por `AK` (grupo).
AQ = 43

DATA_QUALIDADE_OK: Final[str] = 'OK'
INCLUIDO_SIM: Final[str] = 'Sim'
INCLUIDO_NAO: Final[str] = 'Não'
