"""Base de apoio R–AQ compartilhada entre as abas enriquecidas de INMS —
extraída de `excel/inms_1_1/_raw_block.py` quando o INMS 1.2 ganhou seu
próprio renderer enriquecido.

Escreve só as colunas que não dependem de um prazo contratual fixo em horas
corridas (INMS 1.1: 2h para incidentes de criticidade alta) — as colunas
AB/AC/AI (limite contratual bruto e sua ordem de amostragem) e AE (controle
contratual bruto) ficam a cargo de `excel/inms_1_1/_raw_block.py`, chamada
como extensão só por aquela aba: o INMS 1.2 (requisições) tem prazo
variável por SLA embutido em texto livre (`"(CIT) Requisição - Baixa 16
horas"`), então recalcular um limite bruto exigiria parsear essa string —
frágil e fora de escopo (ver docs/spec/inms-pipeline.md §2).
"""

from __future__ import annotations

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as cl
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.categoria_filter import GRUPO_EXECUTOR_COLUMN
from pyauditor.excel._datetime import parse_dt
from pyauditor.excel._inms_audit_common._layout import (
    AA,
    AD,
    AF,
    AG,
    AH,
    AJ,
    AK,
    AL,
    AM,
    AN,
    AO,
    AP,
    AQ,
    ATIVIDADE_COLUMN,
    BODY_FONT,
    DATA_FIM_COLUMN,
    DATA_LIMITE_COLUMN,
    DATA_QUALIDADE_OK,
    DATA_SOLICITACAO_COLUMN,
    DATETIME_FMT,
    DUR,
    HEADER_FILL,
    HEADER_FONT,
    INCLUIDO_SIM,
    NO_PRAZO_COLUMN,
    NOTE_FONT,
    NUM_SOLICITACAO_COLUMN,
    RED_FILL,
    TECNICO_COLUMN,
    R,
    S,
    T,
    U,
    V,
    W,
    X,
    Y,
    Z,
)
from pyauditor.excel._safety import safe_excel_text

_CORE_HEADERS = {
    R: 'Nº Solicitação',
    S: 'Grupo executor',
    T: 'Atividade',
    U: 'DataHoraSolicitacao',
    V: 'DataHoraLimite (ITSM)',
    W: 'DataHoraFim',
    X: 'No prazo (fornecedor)',
    Y: 'TecnicoExecutor',
    Z: 'Nível',
    AA: 'Categoria',
    AD: 'No prazo (data limite ITSM)',
    AF: 'Atraso vs. limite ITSM (min)',
    AG: 'Duração criação→resolução (dias)',
    AH: 'Ordem — fora do prazo',
    AJ: 'Situação dos dados',
    AN: 'Divergência No prazo (fornecedor x ITSM)',
    AO: 'Ordem — divergência fornecedor x ITSM',
    AP: 'Incluído no cálculo (Seção 4)',
    AQ: 'Incluído no INMS? (mapa por grupo)',
}


def write_raw_block_core(
    sheet: Worksheet,
    rows: list[dict[str, str]],
    grupo_rows: list[tuple[str, str, str]],
    last_row: int,
    *,
    first_group_row: int,
) -> None:
    del last_row  # mantido no perfil de chamada por simetria com `raw_range`
    note = sheet.cell(
        row=1,
        column=R - 1,
        value=(
            f'ÁREA DE APOIO — base de dados para as fórmulas desta aba '
            f'({len(rows)} registros do CSV bruto). Não excluir.'
        ),
    )
    note.font = NOTE_FONT
    for col, text in _CORE_HEADERS.items():
        cell = sheet.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[cl(col)].width = 16

    for col, text in {
        AK: 'Grupo executor',
        AL: 'Nível',
        AM: 'Categoria',
    }.items():
        cell = sheet.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        sheet.column_dimensions[cl(col)].width = 30
    for idx, (grupo, nivel, categoria) in enumerate(grupo_rows, start=2):
        sheet.cell(
            row=idx, column=AK, value=safe_excel_text(grupo)
        ).font = BODY_FONT
        sheet.cell(
            row=idx, column=AL, value=safe_excel_text(nivel)
        ).font = BODY_FONT
        sheet.cell(
            row=idx, column=AM, value=safe_excel_text(categoria)
        ).font = BODY_FONT
        # Referência direta (não estruturada) à célula "Incluído no INMS?"
        # da linha do grupo na Seção 4 — grupo_rows alimenta as duas seções
        # na mesma ordem, então a linha correspondente é sempre
        # `first_group_row + offset`.
        section4_row = first_group_row + (idx - 2)
        sheet.cell(
            row=idx, column=AQ, value=f'=I{section4_row}'
        ).font = BODY_FONT
    map_range = f'${cl(AK)}$2:${cl(AM)}${1 + len(grupo_rows)}'
    ak_range = f'${cl(AK)}$2:${cl(AK)}${1 + len(grupo_rows)}'
    aq_range = f'${cl(AQ)}$2:${cl(AQ)}${1 + len(grupo_rows)}'

    for i, row in enumerate(rows, start=2):
        num_solicitacao = safe_excel_text(row[NUM_SOLICITACAO_COLUMN])
        sheet.cell(row=i, column=R, value=num_solicitacao).font = BODY_FONT
        sheet.cell(
            row=i, column=S, value=safe_excel_text(row[GRUPO_EXECUTOR_COLUMN])
        ).font = BODY_FONT
        sheet.cell(
            row=i, column=T, value=safe_excel_text(row[ATIVIDADE_COLUMN])
        ).font = BODY_FONT

        solicitacao = parse_dt(row[DATA_SOLICITACAO_COLUMN])
        limite = parse_dt(row[DATA_LIMITE_COLUMN])
        fim = parse_dt(row[DATA_FIM_COLUMN])

        u_cell = sheet.cell(row=i, column=U, value=solicitacao.value)
        u_cell.number_format = DATETIME_FMT
        v_cell = sheet.cell(row=i, column=V, value=limite.value)
        v_cell.number_format = DATETIME_FMT
        w_cell = sheet.cell(row=i, column=W, value=fim.value)
        w_cell.number_format = DATETIME_FMT
        for c in (u_cell, v_cell, w_cell):
            c.font = BODY_FONT

        if solicitacao.is_malformed:
            qualidade = 'Data de abertura inválida'
        elif solicitacao.is_blank:
            qualidade = 'Data de abertura ausente'
        elif limite.is_malformed:
            qualidade = 'Data limite inválida'
        elif fim.is_malformed:
            qualidade = 'Data de encerramento inválida'
        elif (
            solicitacao.value is not None
            and fim.value is not None
            and fim.value < solicitacao.value
        ):
            qualidade = 'Encerramento anterior à abertura'
        else:
            qualidade = DATA_QUALIDADE_OK
        aj_cell = sheet.cell(row=i, column=AJ, value=qualidade)
        aj_cell.font = BODY_FONT
        if qualidade != DATA_QUALIDADE_OK:
            aj_cell.fill = RED_FILL

        sheet.cell(
            row=i, column=X, value=safe_excel_text(row[NO_PRAZO_COLUMN])
        ).font = BODY_FONT
        sheet.cell(
            row=i, column=Y, value=safe_excel_text(row[TECNICO_COLUMN])
        ).font = BODY_FONT

        sc, uc, vc, wc = (
            f'{cl(S)}{i}',
            f'{cl(U)}{i}',
            f'{cl(V)}{i}',
            f'{cl(W)}{i}',
        )
        z_cell = sheet.cell(
            row=i,
            column=Z,
            value=f'=IFERROR(VLOOKUP({sc},{map_range},2,FALSE),"")',
        )
        aa_cell = sheet.cell(
            row=i,
            column=AA,
            value=f'=IFERROR(VLOOKUP({sc},{map_range},3,FALSE),"")',
        )
        ad_cell = sheet.cell(
            row=i,
            column=AD,
            value=f'=IF(OR({wc}="",{vc}=""),"",IF({wc}<={vc},"S","N"))',
        )
        af_cell = sheet.cell(
            row=i,
            column=AF,
            value=f'=IF(AND({wc}<>"",{vc}<>"",{wc}>{vc}),({wc}-{vc})*1440,0)',
        )
        af_cell.number_format = '0.0'
        ag_cell = sheet.cell(
            row=i,
            column=AG,
            value=f'=IF(OR({uc}="",{wc}="",{wc}<{uc}),"",{wc}-{uc})',
        )
        ag_cell.number_format = DUR
        for c in (z_cell, aa_cell, ad_cell):
            c.font = BODY_FONT

        xc = f'{cl(X)}{i}'
        ah_cell = sheet.cell(
            row=i, column=AH, value=f'=IF({xc}="N",COUNTIF($X$2:{xc},"N"),"")'
        )
        ah_cell.font = BODY_FONT

        adc = f'{cl(AD)}{i}'
        an_cell = sheet.cell(
            row=i,
            column=AN,
            value=f'=IF(OR({xc}="",{adc}=""),"",IF({xc}<>{adc},"Sim","Não"))',
        )
        an_cell.font = BODY_FONT
        anc = f'{cl(AN)}{i}'
        ao_cell = sheet.cell(
            row=i,
            column=AO,
            value=f'=IF({anc}="Sim",COUNTIF($AN$2:{anc},"Sim"),"")',
        )
        ao_cell.font = BODY_FONT

        # Lookup ao vivo contra o toggle "Incluído no INMS?" da Seção 4, via
        # o mapa por grupo `AQ` (que por sua vez referencia a célula da
        # Seção 4 diretamente — ver acima); `IFERROR`/`"Sim"` como padrão
        # cobre o caso (não deveria ocorrer) de um grupo sem linha
        # correspondente no mapa.
        ap_cell = sheet.cell(
            row=i,
            column=AP,
            value=(
                f'=IFERROR(INDEX({aq_range},MATCH({sc},{ak_range},0)),'
                f'"{INCLUIDO_SIM}")'
            ),
        )
        ap_cell.font = BODY_FONT
