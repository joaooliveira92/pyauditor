"""Seções 6–7 da aba INMS 1.3 (projetos fora do prazo, auditoria) — espelha
`excel/inms_1_1/_sections_6_7.py` (mesmo shape `ratio` de ratio único), mas
sem o controle contra um prazo contratual bruto de N horas corridas: o INMS
1.3 (config `configs/_shared/inms-03.yaml`) não declara nenhuma constante
de horas equivalente à do INMS 1.1
(`excel/inms_1_1/_layout.py:_PRAZO_HORAS_CORRIDAS`), então inventar uma
aqui seria uma regra contratual não confirmada — mesma razão de
`excel/inms_1_2/_sections_6_7.py`, cuja Seção 7 simplificada esta segue de
perto (sem a segmentação por categoria SLA do INMS 1.2, já que o INMS 1.3
é um ratio único). A Seção 7b (divergência fornecedor x ITSM) é genérica e
reaproveitada de `excel/_inms_audit_common/_section_7b.py`.
"""

from __future__ import annotations

from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common import _section_7b
from pyauditor.excel._inms_audit_common._cells import (
    ColumnRange,
    add_situacao_conditional_formatting,
    add_table,
    header_row,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    AD,
    AF,
    AH,
    AP,
    BODY_FONT,
    BORDER,
    DATETIME_FMT,
    INCLUIDO_SIM,
    LABEL_FONT,
    NO_PRAZO_COLUMN,
    NOTE_FONT,
    PCT4,
    R,
    RED_FILL,
    S,
    T,
    U,
    UNLOCKED,
    V,
    W,
    Y,
)


def write_section_6_fora_prazo(
    sheet: Worksheet,
    *,
    rows: list[dict[str, str]],
    rng: ColumnRange,
    start_row: int,
    table_name: str,
) -> int:
    """Devolve `next_free_row` — a linha livre após a nota "nenhum projeto"
    ou após a tabela, usada pela Seção 7 como sua própria linha inicial."""
    s6_bar = start_row
    section_bar(
        sheet, s6_bar, 'SEÇÃO 6 · PROJETOS FORA DO PRAZO', last_col=11
    )
    header_row(
        sheet,
        s6_bar + 1,
        (
            'Nº solicitação',
            'Grupo executor',
            'Atividade',
            'Abertura',
            'Limite (ITSM)',
            'Encerramento',
            'Atraso vs. limite (min)',
            'Técnico executor',
            'Justificativa',
            'Aceite da justificativa',
            'Documento/evidência',
        ),
        numeric_cols=frozenset({4, 5, 6, 7}),
    )
    fora_first = s6_bar + 2
    # Tabela dimensionada ao número real de projetos fora do prazo (não um
    # teto arbitrário) — a Seção 6 existe para evidenciar exceções, não pra
    # truncá-las quando o órgão tiver mais de um punhado.
    fora_count = sum(1 for row in rows if row[NO_PRAZO_COLUMN] == 'N')
    ah_range = rng(AH)
    if fora_count == 0:
        sheet.merge_cells(f'A{fora_first}:K{fora_first}')
        sheet[f'A{fora_first}'] = 'Nenhum projeto fora do prazo no período.'
        sheet[f'A{fora_first}'].font = NOTE_FONT
        return fora_first + 2

    for n in range(1, fora_count + 1):
        r = fora_first + n - 1
        match_expr = f'MATCH({n},{ah_range},0)'
        c1 = sheet.cell(
            row=r, column=1, value=f'=IFERROR(INDEX({rng(R)},{match_expr}),"")'
        )
        c2 = sheet.cell(
            row=r, column=2, value=f'=IFERROR(INDEX({rng(S)},{match_expr}),"")'
        )
        c3 = sheet.cell(
            row=r, column=3, value=f'=IFERROR(INDEX({rng(T)},{match_expr}),"")'
        )
        c4 = sheet.cell(
            row=r, column=4, value=f'=IFERROR(INDEX({rng(U)},{match_expr}),"")'
        )
        c4.number_format = DATETIME_FMT
        c5 = sheet.cell(
            row=r, column=5, value=f'=IFERROR(INDEX({rng(V)},{match_expr}),"")'
        )
        c5.number_format = DATETIME_FMT
        c6 = sheet.cell(
            row=r, column=6, value=f'=IFERROR(INDEX({rng(W)},{match_expr}),"")'
        )
        c6.number_format = DATETIME_FMT
        c7 = sheet.cell(
            row=r,
            column=7,
            value=f'=IFERROR(INDEX({rng(AF)},{match_expr}),"")',
        )
        c7.number_format = '0'
        c8 = sheet.cell(
            row=r, column=8, value=f'=IFERROR(INDEX({rng(Y)},{match_expr}),"")'
        )
        c9 = sheet.cell(row=r, column=9, value='Não informado')
        c10 = sheet.cell(row=r, column=10, value='Não informado')
        c11 = sheet.cell(row=r, column=11, value='Não informado')
        for c in (c1, c2, c3, c4, c5, c6, c7, c8, c9, c10, c11):
            c.font = BODY_FONT
            c.border = BORDER
            # Linha inteira sinalizada — não só a 1ª coluna — para deixar
            # claro que o projeto todo está fora do prazo, não só o Nº.
            c.fill = RED_FILL
        # Justificativa/aceite/evidência são preenchimento manual da
        # auditoria — permanecem editáveis com a planilha protegida.
        c9.protection = UNLOCKED
        c10.protection = UNLOCKED
        c11.protection = UNLOCKED
    fora_last = fora_first + fora_count - 1
    add_table(sheet, table_name, f'A{s6_bar + 1}:K{fora_last}')
    return fora_last + 2


def write_section_7_auditoria(
    sheet: Worksheet,
    *,
    rows: list[dict[str, str]],
    rng: ColumnRange,
    start_row: int,
    table_name_fornecedor_itsm: str,
) -> int:
    """Devolve `next_free_row` — a linha livre após a Seção 7b, usada pela
    Seção 8 como sua própria linha inicial. Sem o controle contra o prazo
    contratual bruto (ver docstring do módulo)."""
    s7_bar = start_row
    section_bar(
        sheet, s7_bar, 'SEÇÃO 7 · AUDITORIA DO PRAZO CONTRATUAL', last_col=7
    )
    sheet.cell(
        row=s7_bar + 1, column=1, value='Controles de resultado'
    ).font = LABEL_FONT
    header_row(
        sheet,
        s7_bar + 2,
        ('Metodologia', 'Resultado', 'Situação'),
        numeric_cols=frozenset({2}),
    )
    # `AP` filtra os mesmos grupos excluídos pelo toggle da Seção 4, para
    # que estes controles continuem batendo com B13/C13 (já filtrados).
    ap_sim = f'{rng(AP)},"{INCLUIDO_SIM}"'
    ctrl_rows = [
        ("Resultado informado pelo fornecedor (campo 'No prazo')", 'C13/B13'),
        (
            'Resultado reproduzido pela data limite registrada no ITSM '
            '(DataHoraFim ≤ DataHoraLimite)',
            f'COUNTIFS({rng(AD)},"S",{ap_sim})/B13',
        ),
    ]
    first_ctrl = s7_bar + 3
    for i, (label, division) in enumerate(ctrl_rows):
        r = first_ctrl + i
        sheet.cell(row=r, column=1, value=label).font = BODY_FONT
        sheet.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
        val = sheet.cell(
            row=r, column=2, value=f'=IF(B13=0,"Sem ocorrências",{division})'
        )
        val.number_format = PCT4
        sit_formula = (
            f'=IF(B13=0,"Não '
            f'aplicável",IF(B{r}>=A13,"Meta '
            f'atingida","Meta '
            f'não '
            f'atingida"))'
        )
        sit = sheet.cell(row=r, column=3, value=sit_formula)
        val.font = BODY_FONT
        sit.font = BODY_FONT
        for col in range(1, 4):
            sheet.cell(row=r, column=col).border = BORDER
        add_situacao_conditional_formatting(sheet, sit.coordinate)

    next_free_row = first_ctrl + len(ctrl_rows) + 1
    return _section_7b.write_section_7b_divergencia_fornecedor(
        sheet,
        rows=rows,
        rng=rng,
        start_row=next_free_row,
        table_name=table_name_fornecedor_itsm,
    )
