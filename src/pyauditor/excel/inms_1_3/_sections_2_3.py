"""Seções 2–3 da aba INMS 1.3 (resumo, memória de um único ratio) — espelha
`excel/inms_1_1/_sections_1_3.py` (mesmo shape `ratio`/`count_distinct`),
trocando o vocabulário de "incidentes" por "projetos" (INMS 1.3 = Projetos
atendidos dentro do prazo). A Seção 1 (identificação) é genérica e chamada
diretamente de `excel/_inms_audit_common/_section_1.py` por
`inms_1_3/write.py`, sem wrapper local.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common._cells import (
    add_situacao_conditional_formatting,
    header_row,
    label_value,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    LABEL_FONT,
    NOTE_FONT,
    ORANGE_FILL,
    PCT2,
    PCT4,
    TEAL_FILL,
)


def write_section_2_resumo(
    sheet: Worksheet,
    *,
    iap: str,
    iadp: str,
    fora: str,
    meta_value: float,
) -> int:
    """Seção de linhas fixas (11-14) — devolve `next_free_row` (16, fixo)
    para a Seção 3."""
    section_bar(sheet, 11, 'SEÇÃO 2 · RESUMO EXECUTIVO')
    header_row(
        sheet,
        12,
        (
            'Meta contratual',
            'PAP (abertos)',
            'PDP (dentro do prazo)',
            'Fora do prazo',
            'Resultado INMS 1.3',
            'Diferença p/ meta',
            'Situação',
        ),
    )
    # Cabeçalho recentralizado para casar com o alinhamento centralizado da
    # linha de KPIs abaixo (linha 13) — faixa de indicadores curta, não
    # tabela de dados onde numérico ficaria à direita.
    for col in range(1, 8):
        sheet.cell(row=12, column=col).alignment = Alignment(
            horizontal='center', wrap_text=True, vertical='center'
        )
    sheet['A13'] = f'={meta_value}'
    sheet['B13'] = f'={iap}'
    sheet['C13'] = f'={iadp}'
    sheet['D13'] = f'={fora}'
    # B13=0 (nenhum projeto aberto no período) não é "0% de desempenho" —
    # é resultado não mensurável; guardado para não propagar #DIV/0! e para
    # não afirmar "meta não atingida" indevidamente.
    sheet['E13'] = '=IF(B13=0,"Sem ocorrências",C13/B13)'
    # Meta mínima ("`>=`") — único operador suportado por este renderer.
    sheet['F13'] = '=IF(B13=0,"",E13-A13)'
    sheet['G13'] = (
        '=IF(B13=0,"Sem ocorrências",IF(E13>=A13,"Meta atingida","Meta não '
        'atingida"))'
    )

    sheet['A13'].number_format = PCT2
    sheet['E13'].number_format = PCT2
    sheet['F13'].number_format = PCT2
    for c in ('A13', 'B13', 'C13', 'D13', 'E13', 'F13', 'G13'):
        sheet[c].font = Font(name='Arial', size=12, bold=True)
        sheet[c].fill = TEAL_FILL
        sheet[c].alignment = Alignment(horizontal='center')
    sheet.row_dimensions[13].height = 24

    add_situacao_conditional_formatting(sheet, 'G13')

    sheet.merge_cells('A14:L14')
    pen_note = sheet.cell(
        row=14,
        column=1,
        value='Penalidade: '
        'Pendente '
        'de '
        'confirmação '
        'da '
        'regra '
        'de '
        'arredondamento '
        '— '
        'ver '
        'Seção '
        '9.',
    )
    pen_note.font = LABEL_FONT
    pen_note.fill = ORANGE_FILL
    return 16


def write_section_3_memoria(sheet: Worksheet) -> int:
    """Seção de linhas fixas (16-26) — devolve `next_free_row` (28, fixo)
    para a Seção 4."""
    section_bar(sheet, 16, 'SEÇÃO 3 · MEMÓRIA DO CÁLCULO CONSOLIDADO')
    label_value(
        sheet, 17, 'PAP (projetos abertos no período):', '=B13', fmt='0'
    )
    label_value(sheet, 18, 'PDP (projetos dentro do prazo):', '=C13', fmt='0')
    sheet.merge_cells('A19:C19')
    sheet['A19'] = 'INMS 1.3 = PDP ÷ PAP'
    sheet['A19'].font = NOTE_FONT
    sheet.merge_cells('A20:C20')
    sheet['A20'] = '=CONCATENATE("INMS 1.3 = ",B18," ÷ ",B17)'
    sheet['A20'].font = NOTE_FONT
    label_value(sheet, 21, 'INMS 1.3 (resultado, 4 casas):', '=E13', fmt=PCT4)
    label_value(sheet, 22, 'Meta contratual:', '=A13', fmt=PCT4)
    label_value(
        sheet,
        23,
        'Desvio em pontos percentuais (Resultado - Meta):',
        '=IF(B13=0,"",E13-A13)',
        fmt=PCT4,
    )
    label_value(
        sheet,
        24,
        'Quantidade mínima dentro do prazo p/ atingir a meta:',
        '=ROUNDUP(B13*A13,0)',
        fmt='0',
    )
    label_value(
        sheet,
        25,
        'Margem em quantidade de projetos (PDP - mínimo):',
        '=C13-B24',
        fmt='0;(0)',
    )
    sheet.merge_cells('A26:L26')

    def _narrativa(desfecho: str) -> str:
        return (
            'CONCATENATE("Com ",B17," projetos, seriam necessários pelo '
            'menos ",B24,'
            f'" dentro do prazo para atingir ",TEXT(B22,"0.00%"),". O '
            f'resultado ficou {desfecho})'
        )

    # Narrativa condicional ao sinal da margem (B25 = PDP - mínimo): "abaixo"
    # só quando a margem é negativa — antes disso o texto sempre dizia
    # "abaixo do mínimo" mesmo com meta superada (ABS(B25) escondia o sinal).
    abaixo = _narrativa('",ABS(B25)," projeto(s) abaixo do mínimo necessário."')
    exato = _narrativa('exatamente no mínimo necessário."')
    acima = _narrativa('",B25," projeto(s) acima do mínimo necessário."')
    sheet['A26'] = (
        '=IF(B13=0,"Sem projetos abertos no período — indicador não '
        'mensurável.",'
        f'IF(B25<0,{abaixo},IF(B25=0,{exato},{acima})))'
    )
    sheet['A26'].font = NOTE_FONT
    return 28
