"""Seções 6–7 da aba INMS 1.2 (requisições fora do prazo, auditoria) —
adaptadas do INMS 1.1. A Seção 6 acrescenta a coluna "Prioridade (SLA)" (o
INMS 1.2 segmenta por prioridade, ao contrário do 1.1). A Seção 7 é
simplificada em relação ao 1.1: sem o controle contra um prazo contratual
bruto de N horas corridas — o SLA do INMS 1.2 varia por linha e vem
embutido em texto livre (`SLA`), então recalcular esse limite exigiria
parsear a string (frágil, fora de escopo — ver
`excel/_inms_audit_common/_raw_block.py`). A Seção 7b (divergência
fornecedor x ITSM) é genérica e reaproveitada de
`excel/_inms_audit_common/_section_7b.py`.
"""

from __future__ import annotations

from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common import _section_7b
from pyauditor.excel._inms_audit_common._cells import (
    ColumnRange,
    add_situacao_conditional_formatting,
    add_table,
    header_row,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    AD,
    AF,
    AH,
    AP,
    BODY_FONT,
    BORDER,
    DATETIME_FMT,
    INCLUIDO_SIM,
    LABEL_FONT,
    NO_PRAZO_COLUMN,
    NOTE_FONT,
    PCT4,
    RED_FILL,
    UNLOCKED,
    R,
    S,
    T,
    U,
    V,
    W,
    Y,
)
from pyauditor.excel.inms_1_2._layout import SLA_RAW_COLUMN, CategoryParams


def _write_section_6_fora_prazo(
    sheet: Worksheet,
    *,
    rows: list[dict[str, str]],
    rng: ColumnRange,
    start_row: int,
    table_name: str,
) -> int:
    """Devolve `next_free_row` — a linha livre após a nota "nenhuma
    requisição fora do prazo" ou após a tabela, usada pela Seção 7 como sua
    própria linha inicial."""
    s6_bar = start_row
    section_bar(
        sheet, s6_bar, 'SEÇÃO 6 · REQUISIÇÕES FORA DO PRAZO', last_col=12
    )
    header_row(
        sheet,
        s6_bar + 1,
        (
            'Nº solicitação',
            'Grupo executor',
            'Atividade',
            'Prioridade (SLA)',
            'Abertura',
            'Limite (ITSM)',
            'Encerramento',
            'Atraso vs. limite (min)',
            'Técnico executor',
            'Justificativa',
            'Aceite da justificativa',
            'Documento/evidência',
        ),
        numeric_cols=frozenset({5, 6, 7, 8}),
    )
    fora_first = s6_bar + 2
    # Tabela dimensionada ao número real de requisições fora do prazo (não
    # um teto arbitrário) — a Seção 6 existe para evidenciar exceções, não
    # pra truncá-las quando o órgão tiver mais de um punhado.
    fora_count = sum(1 for row in rows if row[NO_PRAZO_COLUMN] == 'N')
    ah_range = rng(AH)
    if fora_count == 0:
        sheet.merge_cells(f'A{fora_first}:L{fora_first}')
        sheet[f'A{fora_first}'] = 'Nenhuma requisição fora do prazo no período.'
        sheet[f'A{fora_first}'].font = NOTE_FONT
        return fora_first + 2

    for n in range(1, fora_count + 1):
        r = fora_first + n - 1
        match_expr = f'MATCH({n},{ah_range},0)'
        columns = (
            (1, rng(R)),
            (2, rng(S)),
            (3, rng(T)),
            (4, rng(SLA_RAW_COLUMN)),
        )
        cells = [
            sheet.cell(
                row=r,
                column=col,
                value=f'=IFERROR(INDEX({col_range},{match_expr}),"")',
            )
            for col, col_range in columns
        ]
        c5 = sheet.cell(
            row=r, column=5, value=f'=IFERROR(INDEX({rng(U)},{match_expr}),"")'
        )
        c5.number_format = DATETIME_FMT
        c6 = sheet.cell(
            row=r, column=6, value=f'=IFERROR(INDEX({rng(V)},{match_expr}),"")'
        )
        c6.number_format = DATETIME_FMT
        c7 = sheet.cell(
            row=r, column=7, value=f'=IFERROR(INDEX({rng(W)},{match_expr}),"")'
        )
        c7.number_format = DATETIME_FMT
        c8 = sheet.cell(
            row=r,
            column=8,
            value=f'=IFERROR(INDEX({rng(AF)},{match_expr}),"")',
        )
        c8.number_format = '0'
        c9 = sheet.cell(
            row=r, column=9, value=f'=IFERROR(INDEX({rng(Y)},{match_expr}),"")'
        )
        c10 = sheet.cell(row=r, column=10, value='Não informado')
        c11 = sheet.cell(row=r, column=11, value='Não informado')
        c12 = sheet.cell(row=r, column=12, value='Não informado')

        for c in (*cells, c5, c6, c7, c8, c9, c10, c11, c12):
            c.font = BODY_FONT
            c.border = BORDER
            # Linha inteira sinalizada — não só a 1ª coluna — para deixar
            # claro que a requisição toda está fora do prazo, não só o Nº.
            c.fill = RED_FILL
        # Justificativa/aceite/evidência são preenchimento manual da
        # auditoria — permanecem editáveis com a planilha protegida.
        c10.protection = UNLOCKED
        c11.protection = UNLOCKED
        c12.protection = UNLOCKED
    fora_last = fora_first + fora_count - 1
    add_table(sheet, table_name, f'A{s6_bar + 1}:L{fora_last}')
    return fora_last + 2


def _write_section_7_auditoria(
    sheet: Worksheet,
    *,
    rows: list[dict[str, str]],
    categories: list[CategoryParams],
    category_rows: list[int],
    consolidado_row: int,
    rng: ColumnRange,
    start_row: int,
    table_name_fornecedor_itsm: str,
) -> int:
    """Devolve `next_free_row` — a linha livre após a Seção 7b, usada pela
    Seção 8 como sua própria linha inicial. Um par de controles
    (fornecedor x ITSM) por categoria — sem o controle contra o prazo
    contratual bruto (ver docstring do módulo)."""
    s7_bar = start_row
    section_bar(
        sheet, s7_bar, 'SEÇÃO 7 · AUDITORIA DO PRAZO CONTRATUAL', last_col=7
    )
    sheet.cell(
        row=s7_bar + 1, column=1, value='Controles de resultado (por categoria)'
    ).font = LABEL_FONT
    header_row(
        sheet,
        s7_bar + 2,
        ('Metodologia', 'Resultado', 'Situação'),
        numeric_cols=frozenset({2}),
    )
    first_ctrl = s7_bar + 3
    row = first_ctrl
    for category, kpi_row in zip(categories, category_rows, strict=True):
        ap_sim = f'{rng(AP)},"{INCLUIDO_SIM}"'
        sla_match = f'{rng(SLA_RAW_COLUMN)},"*{category.sla_contains}*"'
        ctrl_rows = [
            (
                f'Resultado informado pelo fornecedor — {category.label} '
                f"(campo 'No prazo')",
                f'C{kpi_row}/B{kpi_row}',
            ),
            (
                f'Resultado reproduzido pela data limite registrada no '
                f'ITSM — {category.label} (DataHoraFim ≤ DataHoraLimite)',
                f'COUNTIFS({rng(AD)},"S",{ap_sim},{sla_match})/B{kpi_row}',
            ),
        ]
        for label, division in ctrl_rows:
            sheet.cell(row=row, column=1, value=label).font = BODY_FONT
            sheet.cell(row=row, column=1).alignment = Alignment(wrap_text=True)
            val = sheet.cell(
                row=row,
                column=2,
                value=(f'=IF(B{kpi_row}=0,"Sem ocorrências",{division})'),
            )
            val.number_format = PCT4
            sit = sheet.cell(
                row=row,
                column=3,
                value=(
                    f'=IF(B{kpi_row}=0,"Não aplicável",IF(B{row}>=A'
                    f'{kpi_row},"Meta atingida","Meta não atingida"))'
                ),
            )
            val.font = BODY_FONT
            sit.font = BODY_FONT
            for col in range(1, 4):
                sheet.cell(row=row, column=col).border = BORDER
            add_situacao_conditional_formatting(sheet, sit.coordinate)
            row += 1

    next_free_row = row + 1
    return _section_7b.write_section_7b_divergencia_fornecedor(
        sheet,
        rows=rows,
        rng=rng,
        start_row=next_free_row,
        table_name=table_name_fornecedor_itsm,
        consolidado_row=consolidado_row,
    )
