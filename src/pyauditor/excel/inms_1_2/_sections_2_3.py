"""Seções 2–3 da aba INMS 1.2 (resumo executivo, memória de cálculo) —
adaptadas do INMS 1.1 (`excel/inms_1_1/_sections_1_3.py`) para o shape
`segmented_ratio`: 3 sub-ratios (Alta/Média/Baixa, por `SLA`), cada um com
sua própria linha de KPIs contra a mesma meta contratual, mais uma linha
consolidada (pooled — soma dos numeradores/denominadores das 3 categorias),
igual à semântica de `SegmentedRatioStrategy.calculate` (`result_pct`
pooled, `conforms` só quando as 3 categorias batem a meta individualmente).

A linha consolidada mantém as mesmas colunas B/C/D/E (IAP/IADP/Fora/
Resultado) que o INMS 1.1 grava na linha 13 fixa — as Seções 4/5/7b
compartilhadas (`excel/_inms_audit_common/`) leem essas colunas por
`consolidado_row` (linha variável, não mais fixa em 13).
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common._cells import (
    ColumnRange,
    add_situacao_conditional_formatting,
    header_row,
    label_value,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    AP,
    INCLUIDO_SIM,
    LABEL_FONT,
    NOTE_FONT,
    ORANGE_FILL,
    PCT2,
    PCT4,
    TEAL_FILL,
    X,
)
from pyauditor.excel.inms_1_2._layout import SLA_RAW_COLUMN, CategoryParams

_KPI_HEADERS = (
    'Meta contratual',
    'IAP (abertos)',
    'IADP (dentro do prazo)',
    'Fora do prazo',
    'Resultado',
    'Diferença p/ meta',
    'Situação',
)


def write_section_2_resumo(
    sheet: Worksheet,
    *,
    categories: list[CategoryParams],
    meta_value: float,
    rng: ColumnRange,
    start_row: int,
) -> tuple[int, list[int], int]:
    """Devolve `(next_free_row, category_rows, consolidado_row)` — as
    linhas de KPI de cada categoria (na ordem de `categories`) e a linha
    consolidada (pooled), para a Seção 3 e para as Seções 5/7b montarem
    suas próprias fórmulas/verificações cruzadas."""
    bar_row = start_row
    section_bar(sheet, bar_row, 'SEÇÃO 2 · RESUMO EXECUTIVO')
    header_row(sheet, bar_row + 1, _KPI_HEADERS)
    for col in range(1, 8):
        sheet.cell(row=bar_row + 1, column=col).alignment = Alignment(
            horizontal='center', wrap_text=True, vertical='center'
        )

    category_rows: list[int] = []
    row = bar_row + 2
    for category in categories:
        label_cell = sheet.cell(
            row=row, column=1, value=f'Prioridade {category.label} (SLA)'
        )
        label_cell.font = LABEL_FONT
        row += 1
        kpi_row = row
        category_rows.append(kpi_row)
        sla_range = rng(SLA_RAW_COLUMN)
        ap_sim = f'{rng(AP)},"{INCLUIDO_SIM}"'
        sla_match = f'{sla_range},"*{category.sla_contains}*"'
        iap = f'COUNTIFS({ap_sim},{sla_match})'
        iadp = f'COUNTIFS({ap_sim},{sla_match},{rng(X)},"S")'
        fora = f'COUNTIFS({ap_sim},{sla_match},{rng(X)},"N")'
        _write_kpi_row(
            sheet, kpi_row, meta_value=meta_value, iap=iap, iadp=iadp, fora=fora
        )
        row += 1

    label_cell = sheet.cell(
        row=row, column=1, value='Consolidado (3 categorias — pooled)'
    )
    label_cell.font = LABEL_FONT
    row += 1
    consolidado_row = row
    iap_pool = '+'.join(f'B{r}' for r in category_rows)
    iadp_pool = '+'.join(f'C{r}' for r in category_rows)
    fora_pool = '+'.join(f'D{r}' for r in category_rows)
    situacao_refs = ','.join(f'G{r}="Meta atingida"' for r in category_rows)
    _write_kpi_row(
        sheet,
        consolidado_row,
        meta_value=meta_value,
        iap=iap_pool,
        iadp=iadp_pool,
        fora=fora_pool,
        situacao_override=(
            f'=IF(B{consolidado_row}=0,"Sem ocorrências",'
            f'IF(AND({situacao_refs}),"Meta atingida","Meta não atingida"))'
        ),
    )
    row += 1

    sheet.merge_cells(f'A{row}:L{row}')
    pen_note = sheet.cell(
        row=row,
        column=1,
        value=(
            'Penalidade: soma das 3 categorias — ver Seção 9. Cada '
            'categoria descumprida contribui independentemente ao total.'
        ),
    )
    pen_note.font = LABEL_FONT
    pen_note.fill = ORANGE_FILL
    return row + 2, category_rows, consolidado_row


def _write_kpi_row(
    sheet: Worksheet,
    row: int,
    *,
    meta_value: float,
    iap: str,
    iadp: str,
    fora: str,
    situacao_override: str | None = None,
) -> None:
    sheet.cell(row=row, column=1, value=f'={meta_value}')
    sheet.cell(row=row, column=2, value=f'={iap}')
    sheet.cell(row=row, column=3, value=f'={iadp}')
    sheet.cell(row=row, column=4, value=f'={fora}')
    # B{row}=0 (nenhum registro no período) não é "0% de desempenho" — é
    # resultado não mensurável; guardado para não propagar #DIV/0! e para
    # não afirmar "meta não atingida" indevidamente.
    sheet.cell(
        row=row,
        column=5,
        value=f'=IF(B{row}=0,"Sem ocorrências",C{row}/B{row})',
    )
    sheet.cell(row=row, column=6, value=f'=IF(B{row}=0,"",E{row}-A{row})')
    sheet.cell(
        row=row,
        column=7,
        value=situacao_override
        or (
            f'=IF(B{row}=0,"Sem ocorrências",IF(E{row}>=A{row},'
            f'"Meta atingida","Meta não atingida"))'
        ),
    )
    sheet.cell(row=row, column=1).number_format = PCT2
    sheet.cell(row=row, column=5).number_format = PCT2
    sheet.cell(row=row, column=6).number_format = PCT2
    for col in range(1, 8):
        cell = sheet.cell(row=row, column=col)
        cell.font = Font(name='Arial', size=11, bold=True)
        cell.fill = TEAL_FILL
        cell.alignment = Alignment(horizontal='center')
    sheet.row_dimensions[row].height = 22
    add_situacao_conditional_formatting(sheet, f'G{row}')


def write_section_3_memoria(
    sheet: Worksheet,
    *,
    categories: list[CategoryParams],
    category_rows: list[int],
    consolidado_row: int,
    start_row: int,
) -> int:
    """Devolve `next_free_row` para a Seção 4."""
    bar_row = start_row
    section_bar(sheet, bar_row, 'SEÇÃO 3 · MEMÓRIA DO CÁLCULO CONSOLIDADO')
    row = bar_row + 1
    for category, kpi_row in zip(categories, category_rows, strict=True):
        sheet.cell(
            row=row, column=1, value=f'{category.label} = IADP ÷ IAP'
        ).font = NOTE_FONT
        row += 1
        label_value(
            sheet,
            row,
            f'{category.label} — IAP (abertos):',
            f'=B{kpi_row}',
            fmt='0',
        )
        row += 1
        label_value(
            sheet,
            row,
            f'{category.label} — IADP (dentro do prazo):',
            f'=C{kpi_row}',
            fmt='0',
        )
        row += 1
        label_value(
            sheet,
            row,
            f'{category.label} — Resultado (4 casas):',
            f'=E{kpi_row}',
            fmt=PCT4,
        )
        row += 1
        label_value(
            sheet,
            row,
            f'{category.label} — Quantidade mínima dentro do prazo p/ '
            f'atingir a meta:',
            f'=IF(B{kpi_row}=0,"",ROUNDUP(B{kpi_row}*A{kpi_row},0))',
            fmt='0',
        )
        row += 1

    sheet.merge_cells(f'A{row}:L{row}')
    sheet[f'A{row}'] = (
        'Resultado consolidado (pooled) = soma dos numeradores das 3 '
        'categorias ÷ soma dos denominadores — não a média simples dos 3 '
        'percentuais.'
    )
    sheet[f'A{row}'].font = NOTE_FONT
    row += 1
    label_value(
        sheet,
        row,
        'Resultado consolidado (4 casas):',
        f'=E{consolidado_row}',
        fmt=PCT4,
    )
    row += 2
    return row
