"""Fronteira pública da aba INMS 1.2: validação de params + `write_sheet`
(composição das 9 seções por cursor de linha) + `has_required_columns`.

Espelha `excel/inms_1_1/write.py`, adaptado para o shape `segmented_ratio`
(3 sub-ratios por prioridade `SLA`, cada um com seu próprio `step_points`,
sem `base_points` — ver `config/models.py:SegmentedRatioCalculation` e
`engine/strategies/segmented_ratio.py`).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.utils import get_column_letter as cl

from pyauditor.categoria_filter import GRUPO_EXECUTOR_COLUMN
from pyauditor.config.categorias import (
    CategoriasFile,
    GrupoExecutorMode,
    WholeIndicatorMode,
)
from pyauditor.excel._inms_audit_common._cells import (
    apply_section_outline,
    protect_support_columns,
    raw_range,
)
from pyauditor.excel._inms_audit_common._domain import (
    build_grupo_rows,
    normalize_no_prazo,
)
from pyauditor.excel._inms_audit_common._layout import NO_PRAZO_COLUMN
from pyauditor.excel._inms_audit_common._section_1 import (
    write_section_1_identificacao,
)
from pyauditor.excel._inms_audit_common._section_8 import (
    write_section_8_tempo,
)
from pyauditor.excel._inms_audit_common._sections_4_5 import (
    write_section_4_detalhamento,
    write_section_5_subtotais,
)
from pyauditor.excel._workbook import (
    create_sheet_atomic,
    force_recalc,
    unique_table_name,
)
from pyauditor.excel.inms_1_2._domain import has_required_columns
from pyauditor.excel.inms_1_2._layout import CategoryParams
from pyauditor.excel.inms_1_2._raw_block import _write_raw_block
from pyauditor.excel.inms_1_2._section_9 import _write_section_9_penalidade
from pyauditor.excel.inms_1_2._sections_2_3 import (
    write_section_2_resumo,
    write_section_3_memoria,
)
from pyauditor.excel.inms_1_2._sections_6_7 import (
    _write_section_6_fora_prazo,
    _write_section_7_auditoria,
)
from pyauditor.periodo import PeriodoAfericao

__all__: Final[tuple[str, ...]] = ('has_required_columns', 'write_sheet')


def _validate_write_sheet_params(
    *,
    target_operator: str,
    target_value: float,
    categories: list[CategoryParams],
    step_size_pct: float,
) -> None:
    """Validação na fronteira de `write_sheet()` — mesma razão de
    `inms_1_1/write.py`: a função é pública e aceita `str`/`float`
    irrestritos; o Pydantic (`config/models.py`) já valida esses mesmos
    limites quando os parâmetros vêm de um YAML de config."""
    if target_operator != '>=':
        raise ValueError(
            f'target_operator {target_operator!r} não suportado por este '
            f'renderer — a aba enriquecida do INMS 1.2 assume meta mínima '
            "('>='); para meta-teto ('<='), use o renderer genérico "
            '(_write_grupo_executor_sheet)'
        )
    if not 0 <= target_value <= 100:
        raise ValueError(f'target_value {target_value!r} fora da faixa 0–100')
    if not categories:
        raise ValueError('categories não pode ser vazio')
    for category in categories:
        if category.step_points < 0:
            raise ValueError(
                f'step_points {category.step_points!r} da categoria '
                f'{category.label!r} não pode ser negativo'
            )
    if step_size_pct <= 0:
        raise ValueError(
            f'step_size_pct {step_size_pct!r} deve ser positivo — usado '
            'como divisor nas fórmulas da Seção 9'
        )


def write_sheet(
    workbook: Workbook,
    sheet_name: str,
    *,
    categorias_file: CategoriasFile,
    grupo_executor_entries: list[tuple[str, GrupoExecutorMode]],
    whole_indicator_entries: list[tuple[str, WholeIndicatorMode]],
    fieldnames: list[str],
    rows: list[dict[str, str]],
    target_operator: str,
    target_value: float,
    categories: list[CategoryParams],
    step_size_pct: float,
    contract: str,
    periodo: PeriodoAfericao | None,
    raw_csv_path: Path,
    generated_at: datetime,
) -> None:
    # categorias.yaml não usa whole_indicator para 1.2 hoje.
    del whole_indicator_entries
    del fieldnames  # validado por has_required_columns antes de chamar

    _validate_write_sheet_params(
        target_operator=target_operator,
        target_value=target_value,
        categories=categories,
        step_size_pct=step_size_pct,
    )

    rows = [{**row, NO_PRAZO_COLUMN: normalize_no_prazo(row)} for row in rows]

    with create_sheet_atomic(workbook, sheet_name) as sheet:
        sheet.sheet_view.showGridLines = False
        for col, width in {
            1: 42,
            2: 24,
            3: 16,
            4: 12,
            5: 14,
            6: 12,
            7: 12,
            8: 16,
            9: 14,
            10: 34,
            11: 20,
            12: 40,
        }.items():
            sheet.column_dimensions[cl(col)].width = width

        last_row = 1 + len(rows)
        real_values = {row[GRUPO_EXECUTOR_COLUMN] for row in rows}
        grupo_rows = build_grupo_rows(
            categorias_file, grupo_executor_entries, real_values
        )

        table_names = {
            'grupo_executor': unique_table_name(
                workbook, 'TabelaGrupoExecutor'
            ),
            'fora_do_prazo': unique_table_name(workbook, 'TabelaForaDoPrazo'),
            'amostra_divergencia_fornecedor': unique_table_name(
                workbook, 'TabelaAmostraDivergenciaFornecedor'
            ),
        }

        def rng(col: int) -> str:
            return raw_range(col, last_row)

        meta_value = target_value / 100

        write_section_1_identificacao(
            sheet,
            title='INMS 1.2 – Requisições atendidas dentro do prazo',
            rows=rows,
            contract=contract,
            periodo=periodo,
            raw_csv_path=raw_csv_path,
            generated_at=generated_at,
        )
        # Barras de Seção (linha do rótulo "SEÇÃO N · ..."), usadas ao final
        # para o agrupamento nativo de linhas (`apply_section_outline`).
        section_bars = [3, 11]
        next_row, category_rows, consolidado_row = write_section_2_resumo(
            sheet,
            categories=categories,
            meta_value=meta_value,
            rng=rng,
            start_row=11,
        )
        section_bars.append(next_row)  # Seção 3
        next_row = write_section_3_memoria(
            sheet,
            categories=categories,
            category_rows=category_rows,
            consolidado_row=consolidado_row,
            start_row=next_row,
        )

        first_group_row = next_row + 2
        _write_raw_block(
            sheet, rows, grupo_rows, last_row, first_group_row=first_group_row
        )

        section_bars.append(next_row)  # Seção 4
        next_row = write_section_4_detalhamento(
            sheet,
            grupo_rows=grupo_rows,
            rng=rng,
            start_row=next_row,
            table_name=table_names['grupo_executor'],
        )
        section_bars.append(next_row)  # Seção 5
        next_row = write_section_5_subtotais(
            sheet, rng=rng, start_row=next_row, consolidado_row=consolidado_row
        )
        section_bars.append(next_row)  # Seção 6
        next_row = _write_section_6_fora_prazo(
            sheet,
            rows=rows,
            rng=rng,
            start_row=next_row,
            table_name=table_names['fora_do_prazo'],
        )
        section_bars.append(next_row)  # Seção 7
        next_row = _write_section_7_auditoria(
            sheet,
            rows=rows,
            categories=categories,
            category_rows=category_rows,
            consolidado_row=consolidado_row,
            rng=rng,
            start_row=next_row,
            table_name_fornecedor_itsm=table_names[
                'amostra_divergencia_fornecedor'
            ],
        )
        section_bars.append(next_row)  # Seção 8
        next_row = write_section_8_tempo(sheet, rng=rng, start_row=next_row)
        section_bars.append(next_row)  # Seção 9
        _write_section_9_penalidade(
            sheet,
            categories=categories,
            category_rows=category_rows,
            step_size_pct=step_size_pct,
            start_row=next_row,
        )

        # Cada Seção vira um grupo de linhas colapsável (fim = 2 linhas
        # antes da barra seguinte, pulando a linha em branco separadora; a
        # última Seção vai até `sheet.max_row`).
        section_bounds = []
        for i, bar in enumerate(section_bars):
            if i + 1 < len(section_bars):
                content_end = section_bars[i + 1] - 2
            else:
                content_end = sheet.max_row
            section_bounds.append((bar, content_end))
        apply_section_outline(sheet, section_bounds)

        protect_support_columns(sheet)
        sheet.freeze_panes = 'A2'

        # Impressão/exportação a PDF para auditoria: área restrita às
        # colunas visíveis (A:L — R:AQ são apoio oculto), ajustada à
        # largura da página e com o título (linha 1) repetido em cada
        # página impressa.
        sheet.print_area = f'A1:L{sheet.max_row}'
        sheet.print_title_rows = '1:1'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = (  # ty: ignore[invalid-assignment]
            True
        )

    force_recalc(workbook)
