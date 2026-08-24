"""Base de apoio R–AQ da aba INMS 1.2 — o núcleo compartilhado com o INMS
1.1 vive em `excel/_inms_audit_common/_raw_block.py`; este módulo só
acrescenta o texto bruto da coluna `SLA` (usada para segmentar as fórmulas
por prioridade — Seções 2/9), na coluna `AB` (livre nesta aba)."""

from __future__ import annotations

from openpyxl.utils import get_column_letter as cl
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common._layout import (
    BODY_FONT,
    HEADER_FILL,
    HEADER_FONT,
)
from pyauditor.excel._inms_audit_common._raw_block import write_raw_block_core
from pyauditor.excel._safety import safe_excel_text
from pyauditor.excel.inms_1_2._layout import SLA_COLUMN, SLA_RAW_COLUMN


def _write_raw_block(
    sheet: Worksheet,
    rows: list[dict[str, str]],
    grupo_rows: list[tuple[str, str, str]],
    last_row: int,
    *,
    first_group_row: int,
) -> None:
    write_raw_block_core(
        sheet, rows, grupo_rows, last_row, first_group_row=first_group_row
    )
    _write_sla_column(sheet, rows)


def _write_sla_column(sheet: Worksheet, rows: list[dict[str, str]]) -> None:
    header = sheet.cell(row=1, column=SLA_RAW_COLUMN, value='Prioridade (SLA)')
    header.font = HEADER_FONT
    header.fill = HEADER_FILL
    sheet.column_dimensions[cl(SLA_RAW_COLUMN)].width = 30
    for i, row in enumerate(rows, start=2):
        sheet.cell(
            row=i,
            column=SLA_RAW_COLUMN,
            value=safe_excel_text(row[SLA_COLUMN]),
        ).font = BODY_FONT
