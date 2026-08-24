"""Seção 9 (Penalidade) da aba INMS 1.2 — shape `segmented_ratio`: cada
categoria tem seu próprio `step_points`, sem `base_points` (ao contrário do
INMS 1.1) — replica `SegmentedRatioStrategy.calculate`
(`penalty_i = max(shortfall_i, 0) / step_size_pct * step_points_i`, total =
soma das 3)."""

from __future__ import annotations

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.excel._inms_audit_common._cells import label_value, section_bar
from pyauditor.excel._inms_audit_common._layout import (
    LABEL_FONT,
    NOTE_FONT,
    ORANGE_FILL,
    PCT4,
    TEAL_FILL,
)
from pyauditor.excel.inms_1_2._layout import CategoryParams


def _write_section_9_penalidade(
    sheet: Worksheet,
    *,
    categories: list[CategoryParams],
    category_rows: list[int],
    step_size_pct: float,
    start_row: int,
) -> None:
    s9_bar = start_row
    section_bar(
        sheet, s9_bar, 'SEÇÃO 9 · PENALIDADE (CÁLCULO PRELIMINAR)', last_col=6
    )
    row = s9_bar + 1
    penalty_rows: list[int] = []
    for category, kpi_row in zip(categories, category_rows, strict=True):
        sheet.cell(
            row=row, column=1, value=f'Categoria: {category.label}'
        ).font = LABEL_FONT
        row += 1
        label_value(sheet, row, 'Meta:', f'=A{kpi_row}', fmt=PCT4)
        row += 1
        label_value(sheet, row, 'Resultado:', f'=E{kpi_row}', fmt=PCT4)
        row += 1
        diff_row = row
        label_value(
            sheet,
            diff_row,
            'Diferença (Meta - Resultado):',
            f'=IF(B{kpi_row}=0,"",A{kpi_row}-E{kpi_row})',
            fmt=PCT4,
        )
        row += 1
        diffpp_row = row
        label_value(
            sheet,
            diffpp_row,
            'Diferença em pontos percentuais:',
            f'=IF(B{diff_row}="","",B{diff_row}*100)',
            fmt='0.0000',
        )
        row += 1
        penalty_row = row
        penalty_rows.append(penalty_row)
        below_target = f'E{kpi_row}<A{kpi_row}'
        # Sem `base_points` (diferente do INMS 1.1) — o shape
        # `segmented_ratio` não tem penalidade-base por categoria, só o
        # adicional proporcional (ver `SegmentedRatioStrategy.calculate`).
        label_value(
            sheet,
            penalty_row,
            f'Penalidade — {category.label} ({category.step_points:g} '
            f'pontos a cada {step_size_pct:g} p.p.):',
            (
                f'=IF(B{kpi_row}=0,0,IF({below_target},'
                f'(B{diffpp_row}/{step_size_pct!r})*{category.step_points!r},'
                f'0))'
            ),
            fmt='0.0000',
            fill=TEAL_FILL,
        )
        row += 2

    total_row = row
    total_formula = '+'.join(f'B{r}' for r in penalty_rows)
    label_value(
        sheet,
        total_row,
        'Penalidade total (soma das 3 categorias):',
        f'={total_formula}',
        fmt='0.0000',
        fill=TEAL_FILL,
    )
    sheet.cell(row=total_row, column=1).font = Font(bold=True)
    sheet.cell(row=total_row, column=2).font = Font(bold=True)

    obs_row = total_row + 2
    sheet.merge_cells(f'A{obs_row}:L{obs_row}')
    sheet[f'A{obs_row}'] = (
        'Resultado sujeito à confirmação da regra de arredondamento e das '
        'disposições gerais de glosa do Termo de Referência.'
    )
    sheet[f'A{obs_row}'].font = Font(
        name='Arial', size=10, bold=True, color='9A3412'
    )
    sheet[f'A{obs_row}'].fill = ORANGE_FILL
    sheet[f'A{obs_row}'].alignment = Alignment(wrap_text=True)

    final_note_row = obs_row + 2
    sheet.merge_cells(f'A{final_note_row}:L{final_note_row}')
    sheet[f'A{final_note_row}'] = (
        'Dados de apoio às fórmulas desta aba nas colunas R:AQ (estrutura '
        'de apoio, auditável) — mantidos para rastreabilidade; não '
        'excluir nem reordenar.'
    )
    sheet[f'A{final_note_row}'].font = NOTE_FONT
