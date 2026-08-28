"""Renderer enriquecido para INMS `ratio`/`count_distinct` com
`Grupo_executor` e/ou `whole_indicator` (INMS 1.7, 1.11, 1.12 — satisfação,
telefonia): mesmo padrão de auditoria (identificação, resumo executivo,
detalhe, memória de penalidade) das demais abas enriquecidas do
`sintetico`, mas contando o critério real (`numerator_filter`/
`denominator_filter`, o mesmo predicado de `RatioStrategy._aggregate`) em
vez do "No prazo"/`DataHoraFim` que estas fontes não têm — o mesmo bug de
fundo já corrigido para o INMS 1.14 (ver `precomputed_audit.py`).

A penalidade é calculada uma vez, no nível do indicador (mesma aritmética
de `RatioStrategy._linear_penalty`) — a Seção 3 é uma segmentação
informativa por categoria/grupo, não um cálculo independente por linha (a
oficial nunca segmenta por Grupo_executor). Contagens seguem a convenção
"brutas, pré-quality-gate" do restante do `sintetico` (spec §14.4) — não
usa `accepted_ids`.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as cl
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.categoria_filter import (
    GRUPO_EXECUTOR_COLUMN,
    compute_categoria_values,
)
from pyauditor.config.categorias import (
    CategoriasFile,
    GrupoExecutorMode,
    WholeIndicatorMode,
)
from pyauditor.config.models import RatioCalculation
from pyauditor.engine.strategies import (
    filter_rows,
    meets_target,
    safe_pct,
    shortfall,
)
from pyauditor.excel._inms_audit_common._cells import (
    header_row,
    label_value,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    BODY_FONT,
    NOTE_FONT,
    ORANGE_FILL,
    PCT4,
    TEAL_FILL,
)
from pyauditor.excel._inms_audit_common._section_1 import (
    write_section_1_identificacao,
)
from pyauditor.excel._safety import safe_excel_text
from pyauditor.excel.sintetico._sheets._shared import (
    _NIVEL_BY_CATEGORIA,
    _OUTROS_LABEL,
    _WHOLE_INDICATOR_LABEL,
)
from pyauditor.periodo import PeriodoAfericao

__all__ = ('_write_ratio_audit_sheet',)

_SECTION_3_COLUMNS: Final[tuple[str, ...]] = (
    'Categoria',
    'Nível',
    'Grupo executor',
    'Linhas avaliadas',
    'Atenderam critério',
    'Não atenderam',
    '% resultado',
    'Meta atingida?',
)
_NUMERIC_COLS: Final[frozenset[int]] = frozenset({4, 5, 6, 7})

_COLUMN_WIDTHS: Final[dict[int, int]] = {
    1: 32,
    2: 10,
    3: 34,
    4: 16,
    5: 16,
    6: 16,
    7: 14,
    8: 16,
}

_RESUMO_START_ROW: Final[int] = 11


def _group_stats(
    group_rows: list[dict[str, str]], calculation: RatioCalculation
) -> tuple[int, int, int, float | None]:
    denom_rows = filter_rows(group_rows, calculation.denominator_filter)
    num_rows = filter_rows(denom_rows, calculation.numerator_filter)
    linhas = len(denom_rows)
    atenderam = len(num_rows)
    pct = safe_pct(atenderam, linhas) if linhas else None
    return linhas, atenderam, linhas - atenderam, pct


def _write_resumo_executivo(
    sheet: Worksheet,
    start_row: int,
    *,
    denominator: int,
    numerator: int,
    result_pct: float,
    target_value: float,
    conforms: bool,
) -> int:
    section_bar(sheet, start_row, 'SEÇÃO 2 · RESUMO EXECUTIVO')
    header_row(
        sheet,
        start_row + 1,
        (
            'Linhas avaliadas',
            'Atenderam critério',
            'Resultado',
            'Meta',
            'Situação',
        ),
        numeric_cols=frozenset({1, 2, 3, 4}),
    )
    kpi_row = start_row + 2
    sheet.cell(row=kpi_row, column=1, value=denominator)
    sheet.cell(row=kpi_row, column=2, value=numerator)
    if denominator == 0:
        sheet.cell(row=kpi_row, column=3, value='Sem ocorrências')
        situacao = 'Sem ocorrências no período'
    else:
        cell = sheet.cell(row=kpi_row, column=3, value=result_pct / 100)
        cell.number_format = PCT4
        situacao = 'Alcançado' if conforms else 'Não alcançado'
    meta_cell = sheet.cell(row=kpi_row, column=4, value=target_value / 100)
    meta_cell.number_format = PCT4
    sheet.cell(row=kpi_row, column=5, value=situacao)
    for col in range(1, 6):
        cell = sheet.cell(row=kpi_row, column=col)
        cell.font = Font(name='Arial', size=12, bold=True)
        cell.fill = TEAL_FILL
        cell.alignment = Alignment(horizontal='center')
    sheet.row_dimensions[kpi_row].height = 24
    return kpi_row + 2


def _write_detalhe(
    sheet: Worksheet,
    start_row: int,
    categorias_file: CategoriasFile,
    grupo_executor_entries: list[tuple[str, GrupoExecutorMode]],
    whole_indicator_entries: list[tuple[str, WholeIndicatorMode]],
    calculation: RatioCalculation,
    target_operator: str,
    target_value: float,
    rows: list[dict[str, str]],
) -> int:
    section_bar(sheet, start_row, 'SEÇÃO 3 · DETALHE POR GRUPO EXECUTOR')
    header_row(
        sheet, start_row + 1, _SECTION_3_COLUMNS, numeric_cols=_NUMERIC_COLS
    )
    row_idx = start_row + 2

    real_values = {
        row[GRUPO_EXECUTOR_COLUMN]
        for row in rows
        if GRUPO_EXECUTOR_COLUMN in row
    }
    per_categoria_values, outros_values = (
        compute_categoria_values(grupo_executor_entries, real_values)
        if grupo_executor_entries
        else ({}, set())
    )

    def emit(
        categoria_label: str,
        nivel: str | None,
        grupo: str,
        group_rows: list[dict[str, str]],
    ) -> None:
        nonlocal row_idx
        linhas, atenderam, nao, pct = _group_stats(group_rows, calculation)
        sheet.cell(
            row=row_idx, column=1, value=categoria_label
        ).font = BODY_FONT
        sheet.cell(row=row_idx, column=2, value=nivel or '').font = BODY_FONT
        sheet.cell(
            row=row_idx, column=3, value=safe_excel_text(grupo)
        ).font = BODY_FONT
        sheet.cell(row=row_idx, column=4, value=linhas).font = BODY_FONT
        sheet.cell(row=row_idx, column=5, value=atenderam).font = BODY_FONT
        sheet.cell(row=row_idx, column=6, value=nao).font = BODY_FONT
        if pct is None:
            sheet.cell(row=row_idx, column=7, value='—').font = BODY_FONT
            sheet.cell(row=row_idx, column=8, value='—').font = BODY_FONT
        else:
            cell = sheet.cell(row=row_idx, column=7, value=pct / 100)
            cell.number_format = PCT4
            cell.font = BODY_FONT
            atingiu = meets_target(pct, target_operator, target_value)
            sheet.cell(
                row=row_idx, column=8, value='Sim' if atingiu else 'Não'
            ).font = BODY_FONT
        row_idx += 1

    # ⚡ Bolt: otimização de performance.
    # Pré-agrupa as linhas por Grupo_executor em O(N) para evitar
    # filtragens O(N) repetidas em loops para cada grupo.
    rows_by_grupo: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grupo_val = row.get(GRUPO_EXECUTOR_COLUMN)
        if grupo_val is not None:
            rows_by_grupo[grupo_val].append(row)

    for categoria_key, effective_values in per_categoria_values.items():
        categoria = categorias_file.categorias[categoria_key]
        nivel = _NIVEL_BY_CATEGORIA.get(categoria_key)
        for grupo in sorted(effective_values):
            group_rows = rows_by_grupo.get(grupo, [])
            emit(categoria.label, nivel, grupo, group_rows)

    for categoria_key, _entry in whole_indicator_entries:
        categoria = categorias_file.categorias[categoria_key]
        nivel = _NIVEL_BY_CATEGORIA.get(categoria_key)
        emit(categoria.label, nivel, _WHOLE_INDICATOR_LABEL, rows)

    for grupo in sorted(outros_values):
        group_rows = rows_by_grupo.get(grupo, [])
        emit(_OUTROS_LABEL, None, grupo, group_rows)

    return row_idx


def _write_memoria_penalidade(
    sheet: Worksheet,
    start_row: int,
    *,
    target_operator: str,
    target_value: float,
    result_pct: float,
    penalty_base_points: float,
    penalty_step_points: float,
    penalty_step_size_pct: float,
    penalty: float,
    conforms: bool,
    denominator: int,
) -> int:
    section_bar(
        sheet, start_row, 'SEÇÃO 4 · MEMÓRIA DA PENALIDADE (CÁLCULO INDICADOR)'
    )
    row = start_row + 1
    label_value(sheet, row, 'Meta:', target_value / 100, fmt=PCT4)
    row += 1
    if denominator == 0:
        label_value(
            sheet,
            row,
            'Situação:',
            'Sem ocorrências no período — indicador não mensurável.',
        )
        return row + 1
    label_value(sheet, row, 'Resultado:', result_pct / 100, fmt=PCT4)
    row += 1
    diff = shortfall(result_pct, target_operator, target_value)
    label_value(
        sheet, row, 'Diferença em pontos percentuais:', diff, fmt='0.0000'
    )
    row += 1
    label_value(
        sheet,
        row,
        'Penalidade-base:',
        0 if conforms else penalty_base_points,
        fmt='0',
    )
    row += 1
    additional = (
        0.0
        if conforms
        else max(diff, 0.0) / penalty_step_size_pct * penalty_step_points
    )
    label_value(
        sheet,
        row,
        f'Adicional proporcional ({penalty_step_points:g} pontos a cada '
        f'{penalty_step_size_pct:g} p.p.):',
        additional,
        fmt='0.0000',
    )
    row += 1
    label_value(
        sheet,
        row,
        'Total (base + adicional):',
        penalty,
        fmt='0.0000',
        fill=TEAL_FILL,
    )
    return row + 1


def _write_note(sheet: Worksheet, row: int) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    cell = sheet.cell(
        row=row,
        column=1,
        value=(
            'Contagens da Seção 3 são brutas (pré-quality-gate), para '
            'conferência rápida — não substituem o ROM oficial. A '
            'penalidade (Seção 4) é sempre calculada uma vez, no nível do '
            'indicador inteiro; a segmentação por categoria/grupo executor '
            'é informativa e não gera penalidade própria.'
        ),
    )
    cell.font = NOTE_FONT
    cell.fill = ORANGE_FILL
    cell.alignment = Alignment(wrap_text=True)


def _write_ratio_audit_sheet(
    workbook: Workbook,
    sheet_name: str,
    categorias_file: CategoriasFile,
    grupo_executor_entries: list[tuple[str, GrupoExecutorMode]],
    whole_indicator_entries: list[tuple[str, WholeIndicatorMode]],
    calculation: RatioCalculation,
    target_operator: str,
    target_value: float,
    penalty_base_points: float,
    penalty_step_points: float,
    penalty_step_size_pct: float,
    indicator_name: str,
    rows: list[dict[str, str]],
    contract: str,
    periodo: PeriodoAfericao | None,
    raw_csv_path: Path,
    generated_at: datetime,
) -> None:
    sheet = workbook.create_sheet(title=sheet_name)
    sheet.sheet_view.showGridLines = False
    for col, width in _COLUMN_WIDTHS.items():
        sheet.column_dimensions[cl(col)].width = width

    write_section_1_identificacao(
        sheet,
        title=f'{sheet_name} – {safe_excel_text(indicator_name)}',
        rows=rows,
        contract=contract,
        periodo=periodo,
        raw_csv_path=raw_csv_path,
        generated_at=generated_at,
    )

    denominator_rows = filter_rows(rows, calculation.denominator_filter)
    numerator_rows = filter_rows(denominator_rows, calculation.numerator_filter)
    denominator = len(denominator_rows)
    numerator = len(numerator_rows)
    result_pct = safe_pct(numerator, denominator)

    if denominator == 0:
        conforms = True
        penalty = 0.0
    else:
        conforms = meets_target(result_pct, target_operator, target_value)
        if conforms:
            penalty = 0.0
        else:
            steps = (
                max(shortfall(result_pct, target_operator, target_value), 0.0)
                / penalty_step_size_pct
            )
            penalty = penalty_base_points + steps * penalty_step_points

    next_row = _write_resumo_executivo(
        sheet,
        _RESUMO_START_ROW,
        denominator=denominator,
        numerator=numerator,
        result_pct=result_pct,
        target_value=target_value,
        conforms=conforms,
    )
    next_row = _write_detalhe(
        sheet,
        next_row,
        categorias_file,
        grupo_executor_entries,
        whole_indicator_entries,
        calculation,
        target_operator,
        target_value,
        rows,
    )
    next_row = _write_memoria_penalidade(
        sheet,
        next_row + 1,
        target_operator=target_operator,
        target_value=target_value,
        result_pct=result_pct,
        penalty_base_points=penalty_base_points,
        penalty_step_points=penalty_step_points,
        penalty_step_size_pct=penalty_step_size_pct,
        penalty=penalty,
        conforms=conforms,
        denominator=denominator,
    )
    _write_note(sheet, next_row + 1)
    sheet.freeze_panes = 'A2'
