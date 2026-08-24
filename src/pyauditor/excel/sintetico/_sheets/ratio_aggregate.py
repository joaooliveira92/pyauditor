"""Renderer enriquecido ratio/sum com subtract (INMS 1.6) — mesmo padrão de
auditoria (identificação, resumo executivo, detalhe, memória de penalidade)
das demais abas enriquecidas do `sintetico`. A penalidade é calculada uma
vez, no nível do indicador (mesma aritmética de `RatioStrategy._aggregate`
para `sum`: `(ΣnumeradorColumn - ΣsubtractColumn) / ΣnumeradorColumn`) — a
Seção 3 segmenta por valor de `denominator_filter` (ex. "Acordo de Nível de
Serviço") como detalhe informativo, não como um cálculo independente.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as cl
from openpyxl.worksheet.worksheet import Worksheet

from pyauditor.config.categorias import CategoriasFile, WholeIndicatorMode
from pyauditor.config.models import RatioCalculation
from pyauditor.engine.strategies import (
    filter_rows,
    meets_target,
    parse_decimal,
    safe_pct,
    shortfall,
)
from pyauditor.excel._inms_audit_common._cells import (
    header_row,
    label_value,
    section_bar,
)
from pyauditor.excel._inms_audit_common._layout import (
    BODY_FONT,
    NOTE_FONT,
    ORANGE_FILL,
    PCT4,
    TEAL_FILL,
)
from pyauditor.excel._inms_audit_common._section_1 import (
    write_section_1_identificacao,
)
from pyauditor.excel._safety import safe_excel_text
from pyauditor.excel.sintetico._sheets._shared import _NIVEL_BY_CATEGORIA
from pyauditor.periodo import PeriodoAfericao

__all__ = ('_write_ratio_aggregate_sheet',)

_NUMERIC_COLS: Final[frozenset[int]] = frozenset({4, 5, 6})
_COLUMN_WIDTHS: Final[dict[int, int]] = {
    1: 32,
    2: 10,
    3: 36,
    4: 16,
    5: 16,
    6: 14,
    7: 16,
}
_RESUMO_START_ROW: Final[int] = 11


def _write_resumo_executivo(
    sheet: Worksheet,
    start_row: int,
    *,
    denominator: float,
    numerator: float,
    result_pct: float,
    target_value: float,
    conforms: bool,
) -> int:
    section_bar(sheet, start_row, 'SEÇÃO 2 · RESUMO EXECUTIVO')
    header_row(
        sheet,
        start_row + 1,
        (
            'Total de chamados',
            'Não reabertos',
            'Resultado',
            'Meta',
            'Situação',
        ),
        numeric_cols=frozenset({1, 2, 3, 4}),
    )
    kpi_row = start_row + 2
    sheet.cell(row=kpi_row, column=1, value=int(denominator))
    sheet.cell(row=kpi_row, column=2, value=int(numerator))
    if denominator == 0:
        sheet.cell(row=kpi_row, column=3, value='Sem ocorrências')
        situacao = 'Sem ocorrências no período'
    else:
        cell = sheet.cell(row=kpi_row, column=3, value=result_pct / 100)
        cell.number_format = PCT4
        situacao = 'Alcançado' if conforms else 'Não alcançado'
    meta_cell = sheet.cell(row=kpi_row, column=4, value=target_value / 100)
    meta_cell.number_format = PCT4
    sheet.cell(row=kpi_row, column=5, value=situacao)
    for col in range(1, 6):
        cell = sheet.cell(row=kpi_row, column=col)
        cell.font = Font(name='Arial', size=12, bold=True)
        cell.fill = TEAL_FILL
        cell.alignment = Alignment(horizontal='center')
    sheet.row_dimensions[kpi_row].height = 24
    return kpi_row + 2


def _write_memoria_penalidade(
    sheet: Worksheet,
    start_row: int,
    *,
    target_operator: str,
    target_value: float,
    result_pct: float,
    penalty_base_points: float,
    penalty_step_points: float,
    penalty_step_size_pct: float,
    penalty: float,
    conforms: bool,
    denominator: float,
) -> int:
    section_bar(
        sheet, start_row, 'SEÇÃO 4 · MEMÓRIA DA PENALIDADE (CÁLCULO INDICADOR)'
    )
    row = start_row + 1
    label_value(sheet, row, 'Meta:', target_value / 100, fmt=PCT4)
    row += 1
    if denominator == 0:
        label_value(
            sheet,
            row,
            'Situação:',
            'Sem ocorrências no período — indicador não mensurável.',
        )
        return row + 1
    label_value(sheet, row, 'Resultado:', result_pct / 100, fmt=PCT4)
    row += 1
    diff = shortfall(result_pct, target_operator, target_value)
    label_value(
        sheet, row, 'Diferença em pontos percentuais:', diff, fmt='0.0000'
    )
    row += 1
    label_value(
        sheet,
        row,
        'Penalidade-base:',
        0 if conforms else penalty_base_points,
        fmt='0',
    )
    row += 1
    additional = (
        0.0
        if conforms
        else max(diff, 0.0) / penalty_step_size_pct * penalty_step_points
    )
    label_value(
        sheet,
        row,
        f'Adicional proporcional ({penalty_step_points:g} pontos a cada '
        f'{penalty_step_size_pct:g} p.p.):',
        additional,
        fmt='0.0000',
    )
    row += 1
    label_value(
        sheet,
        row,
        'Total (base + adicional):',
        penalty,
        fmt='0.0000',
        fill=TEAL_FILL,
    )
    return row + 1


def _write_note(sheet: Worksheet, row: int) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    cell = sheet.cell(
        row=row,
        column=1,
        value=(
            'A penalidade (Seção 4) é sempre calculada uma vez, no nível '
            'do indicador inteiro; a segmentação por categoria/acordo de '
            'nível de serviço (Seção 3) é informativa e não gera '
            'penalidade própria.'
        ),
    )
    cell.font = NOTE_FONT
    cell.fill = ORANGE_FILL
    cell.alignment = Alignment(wrap_text=True)


def _write_ratio_aggregate_sheet(
    workbook: Workbook,
    sheet_name: str,
    categorias_file: CategoriasFile,
    entries: list[tuple[str, WholeIndicatorMode]],
    calculation: RatioCalculation,
    target_operator: str,
    target_value: float,
    penalty_base_points: float,
    penalty_step_points: float,
    penalty_step_size_pct: float,
    indicator_name: str,
    rows: list[dict[str, str]],
    contract: str,
    periodo: PeriodoAfericao | None,
    raw_csv_path: Path,
    generated_at: datetime,
) -> None:
    """INMS 1.6 (e qualquer futuro `ratio`/`sum` com
    `sum_numerator_subtract_column`): uma linha por valor distinto da coluna
    de `denominator_filter` (ex. "Acordo de Nível de Serviço"), agregando os
    CSVs elegíveis por grupo com a mesma aritmética de
    `RatioStrategy._aggregate`, em vez do colapso "(indicador inteiro)"."""
    if (
        calculation.denominator_filter is None
        or calculation.sum_numerator_column is None
        or calculation.sum_numerator_subtract_column is None
    ):
        raise ValueError(
            'ratio_aggregate exige `denominator_filter`, '
            '`sum_numerator_column` e `sum_numerator_subtract_column`'
        )
    group_column = calculation.denominator_filter.column
    numerator_column = calculation.sum_numerator_column
    subtract_column = calculation.sum_numerator_subtract_column

    sheet = workbook.create_sheet(title=sheet_name)
    sheet.sheet_view.showGridLines = False
    for col, width in _COLUMN_WIDTHS.items():
        sheet.column_dimensions[cl(col)].width = width

    write_section_1_identificacao(
        sheet,
        title=f'{sheet_name} – {safe_excel_text(indicator_name)}',
        rows=rows,
        contract=contract,
        periodo=periodo,
        raw_csv_path=raw_csv_path,
        generated_at=generated_at,
    )

    eligible_rows = filter_rows(rows, calculation.denominator_filter)
    total_raw = sum(
        parse_decimal(row.get(numerator_column, '') or '0')
        for row in eligible_rows
    )
    total_subtract = sum(
        parse_decimal(row.get(subtract_column, '') or '0')
        for row in eligible_rows
    )
    numerator = total_raw - total_subtract
    denominator = total_raw
    result_pct = safe_pct(numerator, denominator)

    if denominator == 0:
        conforms = True
        penalty = 0.0
    else:
        conforms = meets_target(result_pct, target_operator, target_value)
        if conforms:
            penalty = 0.0
        else:
            steps = (
                max(shortfall(result_pct, target_operator, target_value), 0.0)
                / penalty_step_size_pct
            )
            penalty = penalty_base_points + steps * penalty_step_points

    next_row = _write_resumo_executivo(
        sheet,
        _RESUMO_START_ROW,
        denominator=denominator,
        numerator=numerator,
        result_pct=result_pct,
        target_value=target_value,
        conforms=conforms,
    )

    section_bar(sheet, next_row, 'SEÇÃO 3 · DETALHE POR GRUPO')
    header_row(
        sheet,
        next_row + 1,
        (
            'Categoria',
            'Nível',
            group_column,
            numerator_column,
            subtract_column,
            '% resultado',
            'Meta atingida?',
        ),
        numeric_cols=_NUMERIC_COLS,
    )
    row_idx = next_row + 2

    grupos = list(
        dict.fromkeys(
            row[group_column] for row in eligible_rows if row.get(group_column)
        )
    )

    for categoria_key, _entry in entries:
        categoria = categorias_file.categorias[categoria_key]
        nivel = _NIVEL_BY_CATEGORIA.get(categoria_key)
        for grupo in grupos:
            grupo_rows = [
                row for row in eligible_rows if row[group_column] == grupo
            ]
            grupo_total = sum(
                parse_decimal(row.get(numerator_column, '') or '0')
                for row in grupo_rows
            )
            grupo_subtraido = sum(
                parse_decimal(row.get(subtract_column, '') or '0')
                for row in grupo_rows
            )
            grupo_pct = safe_pct(grupo_total - grupo_subtraido, grupo_total)
            atingiu = (
                meets_target(grupo_pct, target_operator, target_value)
                if grupo_total
                else None
            )
            sheet.cell(
                row=row_idx, column=1, value=categoria.label
            ).font = BODY_FONT
            sheet.cell(
                row=row_idx, column=2, value=nivel or ''
            ).font = BODY_FONT
            sheet.cell(
                row=row_idx, column=3, value=safe_excel_text(grupo)
            ).font = BODY_FONT
            sheet.cell(
                row=row_idx, column=4, value=int(grupo_total)
            ).font = BODY_FONT
            sheet.cell(
                row=row_idx, column=5, value=int(grupo_subtraido)
            ).font = BODY_FONT
            if atingiu is None:
                sheet.cell(row=row_idx, column=6, value='—').font = BODY_FONT
                sheet.cell(row=row_idx, column=7, value='—').font = BODY_FONT
            else:
                cell = sheet.cell(
                    row=row_idx, column=6, value=grupo_pct / 100
                )
                cell.number_format = PCT4
                cell.font = BODY_FONT
                sheet.cell(
                    row=row_idx,
                    column=7,
                    value='Sim' if atingiu else 'Não',
                ).font = BODY_FONT
            row_idx += 1

    next_row = _write_memoria_penalidade(
        sheet,
        row_idx + 1,
        target_operator=target_operator,
        target_value=target_value,
        result_pct=result_pct,
        penalty_base_points=penalty_base_points,
        penalty_step_points=penalty_step_points,
        penalty_step_size_pct=penalty_step_size_pct,
        penalty=penalty,
        conforms=conforms,
        denominator=denominator,
    )
    _write_note(sheet, next_row + 1)
    sheet.freeze_panes = 'A2'
